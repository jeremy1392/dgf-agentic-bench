from __future__ import annotations
import csv, json, os, textwrap
from pathlib import Path
from typing import Any, Dict, List
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from azure_architecture import render_architecture
from evidence_graph import visible_value, node

GATE_TITLES={
    "general":"General Gate — Strategic & Budget Arbitration",
    "it":"IT Gate — Technology Governance & Standards",
    "architecture":"Architecture Gate — Urbanization & Network Constraints",
    "security":"Security Architecture Gate — Cybersecurity & Hardening",
    "tech_readiness":"Tech Readiness Gate — Maturity & Production Reliability",
    "procurement":"Procurement Gate — Purchasing, Contracts & Suppliers",
    "legal":"Legal Gate — Law & Contractual Commitments",
    "compliance":"Compliance Gate — Regulation, Standards & Ethics",
}

GATE_OBJECTIVES={
    "general":"Consolidate specialist opinions and decide whether resources and the next phase can be committed.",
    "it":"Assess technology standards, catalog fit, operability, support, lifecycle, capacity and ITSM readiness.",
    "architecture":"Assess HLD/LLD coherence, APIs, flows, data ownership, network constraints, performance, availability and reversibility.",
    "security":"Assess cybersecurity risk, segmentation, IAM, encryption, hardening, detection, testing and residual risk.",
    "tech_readiness":"Assess test maturity, resilience, backups, recovery, observability, operational handover and aptitude for production.",
    "procurement":"Assess make/buy, tender quality, supplier scoring, TCO and third-party due diligence.",
    "legal":"Assess contractual conformity including liability, SLA, reversibility, IP, data clauses, insurance and signing powers.",
    "compliance":"Assess applicable regulations and standards, privacy, controls, ethics, accessibility, retention and audit evidence.",
}


def _ensure(path:Path): path.parent.mkdir(parents=True,exist_ok=True); return path

def write_json(path:Path,obj): _ensure(path).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding="utf-8")

def write_text(path:Path,text): _ensure(path).write_text(text,encoding="utf-8")

def write_csv(path:Path, rows:List[Dict[str,Any]], fieldnames=None):
    _ensure(path)
    fieldnames=fieldnames or sorted({k for r in rows for k in r})
    with path.open("w",newline="",encoding="utf-8") as f:
        w=csv.DictWriter(f,fieldnames=fieldnames); w.writeheader(); w.writerows(rows)


def _shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def _set_cell_text(cell,text,bold=False):
    cell.text=""; p=cell.paragraphs[0]; r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(9); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def _doc(title,subtitle=None):
    d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    st=d.styles['Normal']; st.font.name='Aptos'; st.font.size=Pt(9.5)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; r=p.add_run(title); r.bold=True; r.font.size=Pt(20)
    if subtitle:
        p=d.add_paragraph(); r=p.add_run(subtitle); r.font.size=Pt(10); r.italic=True
    return d

def _kv(d,title,items):
    d.add_heading(title,level=1); t=d.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    _set_cell_text(t.rows[0].cells[0],'Field',True); _set_cell_text(t.rows[0].cells[1],'Value',True); _shade(t.rows[0].cells[0],'D9EAF7'); _shade(t.rows[0].cells[1],'D9EAF7')
    for k,v in items:
        c=t.add_row().cells; _set_cell_text(c[0],k,True); _set_cell_text(c[1],v)

def _bullets(d,title,items):
    d.add_heading(title,level=1)
    for x in items: d.add_paragraph(str(x),style='List Bullet')

def _save(d,path): _ensure(path); d.save(path); return path


def make_review_requests(case_dir, case, graph, occurrences):
    by_gate={}
    for n in graph["nodes"]:
        for g in n["consumers"]: by_gate.setdefault(g,[]).append(n)
    for gate,nodes in by_gate.items():
        path=case_dir/f"gate_evidence/{gate}/00_Review_Request.docx"
        d=_doc(GATE_TITLES[gate],f"Synthetic review dossier — {case['project']['project_id']}")
        _kv(d,"Project context",[("Project",case['project']['project_name']),("Business unit",case['project']['business_unit']),("Criticality",case['project']['business_criticality']),("Data classification",case['project']['data_classification']),("Target go-live",case['project']['target_go_live'])])
        d.add_heading('Gate objective',level=1); d.add_paragraph(GATE_OBJECTIVES[gate])
        d.add_heading('Evidence package',level=1)
        t=d.add_table(rows=1,cols=4); t.style='Table Grid'
        for i,h in enumerate(['Evidence ID','Version','Status','Path']): _set_cell_text(t.rows[0].cells[i],h,True); _shade(t.rows[0].cells[i],'D9EAF7')
        for n in nodes:
            c=t.add_row().cells
            vals=[n['evidence_id'],n['version'],n['public_status'],n['path']]
            for i,v in enumerate(vals): _set_cell_text(c[i],v)
        _bullets(d,'Synthetic stakeholder comments',[
            f"Project Manager: Please avoid delaying the {case['project']['target_go_live']} target unless the issue is a genuine blocker.",
            f"Sponsor: The project is priority {case['general']['portfolio_priority']} and should proceed where risk can be bounded.",
            "Evidence Owner: Some documents may be stale or incomplete; use authoritative systems where available.",
        ])
        _save(d,path)


def emit_all(case_dir:Path, case:Dict[str,Any], graph:Dict[str,Any], occurrences, reference_results):
    case_dir.mkdir(parents=True,exist_ok=True)
    p=case["project"]
    def v(eid,key,truth): return visible_value(case,graph,eid,key,truth)
    # Shared project charter
    d=_doc('Project Charter',f"Synthetic DGF-Bench case {p['project_id']}")
    _kv(d,'Project',[("Name",p['project_name']),("Route",p['route']),("Business unit",p['business_unit']),("Sponsor",p['sponsor']),("Project manager",p['project_manager']),("Business owner",p['business_owner']),("Risk owner",p['risk_owner']),("Criticality",p['business_criticality']),("Data classification",p['data_classification']),("Users impacted",p['users_impacted']),("Target go-live",p['target_go_live'])])
    _bullets(d,'Declared project assumptions',["Architecture diagrams and supplier documents are supplied by their respective owners.","Gate reviewers must request missing evidence rather than infer it.","A negative or rework decision is a valid outcome when the contract is not satisfied."])
    _save(d,case_dir/'shared/project_charter.docx')

    # GENERAL
    g=case['general']
    write_csv(case_dir/'gate_evidence/general/business_case.csv',[{"metric":"ROI","value":v('BUSINESS_CASE','roi',g['roi'])},{"metric":"NPV_EUR","value":v('BUSINESS_CASE','npv',g['npv_eur'])},{"metric":"Payback_months","value":v('BUSINESS_CASE','payback',g['payback_months'])},{"metric":"Strategic_alignment","value":v('BUSINESS_CASE','strategic_alignment',g['strategic_alignment'])},{"metric":"Portfolio_priority","value":g['portfolio_priority']}])
    write_csv(case_dir/'gate_evidence/general/budget_approval.csv',[{"requested_eur":g['budget_requested_eur'],"approved_eur":g['budget_approved_eur'],"approver":"CFO Office","status":"APPROVED" if g['budget_approved_eur']>=g['budget_requested_eur'] else "PARTIAL"}])
    d=_doc('Benefits Realization Plan'); _kv(d,'Benefits',[("KPI",v('BENEFITS_PLAN','kpi',g['benefits_kpi'])),("Owner",g['benefits_owner']),("Change plan",v('BENEFITS_PLAN','change_plan_status',g['change_plan_status']))]); _save(d,case_dir/'gate_evidence/general/benefits_plan.docx')
    write_csv(case_dir/'gate_evidence/general/portfolio_snapshot.csv',[{"project_id":p['project_id'],"priority":g['portfolio_priority'],"strategic_alignment":g['strategic_alignment'],"requested_eur":g['budget_requested_eur'],"approved_eur":g['budget_approved_eur'],"target_go_live":p['target_go_live']}])
    d=_doc('Governance Committee Briefing',f"Pre-read for {p['project_name']}")
    _kv(d,'Decision context',[("Portfolio priority",g['portfolio_priority']),("Strategic alignment",v('COMMITTEE_BRIEFING','strategic_alignment',g['strategic_alignment'])),("Requested budget",g['budget_requested_eur']),("Approved budget",g['budget_approved_eur']),("ROI",v('COMMITTEE_BRIEFING','roi',g['roi'])),("Risk owner",p['risk_owner'])])
    _bullets(d,'Committee questions',["Are all specialist gate reservations resolved or explicitly owned?","Is the requested funding within the approved envelope?","Are residual risks assigned to a competent business risk owner?","Are benefits and post-go-live tracking owned?"])
    _save(d,case_dir/'gate_evidence/general/committee_briefing.docx')

    # IT
    it=case['it']
    write_csv(case_dir/'gate_evidence/it/cmdb_export.csv',[{"application":p['project_name'],"record_present":it['cmdb_record_present'],"run_owner":p['service_owner'] if it['run_owner_present'] else "","support_model":it['support_model'],"continuity_class":it['service_continuity_class']}])
    write_json(case_dir/'gate_evidence/it/technology_catalog.json',{"catalog_status":it['catalog_status'],"waiver_present":it['waiver_present'],"waiver_expiry_days":it['waiver_expiry_days'],"hosting_model":it['hosting_model']})
    write_csv(case_dir/'gate_evidence/it/license_position.csv',[{"license_compliant":it['license_compliant'],"annual_cost_eur":it['annual_license_cost_eur']}])
    write_csv(case_dir/'gate_evidence/it/lifecycle_eol.csv',[{"technology":it['hosting_model'],"eol_months":it['technology_eol_months'],"capacity_headroom_pct":it['capacity_headroom_pct']}])
    write_csv(case_dir/'gate_evidence/it/itsm_change.csv',[{"change_id":f"CHG{case['seed']%100000:05d}","status":it['change_record_status'],"planned_go_live":p['target_go_live']}])
    d=_doc('Support & Operations RACI'); _kv(d,'Operating model',[("Run owner",p['service_owner'] if v('SUPPORT_RACI','run_owner',it['run_owner_present']) else 'TBD'),("Support model",it['support_model']),("Patch SLA days",it['patch_sla_days']),("CMDB required",'Yes')]); _save(d,case_dir/'gate_evidence/it/support_raci.docx')
    write_csv(case_dir/'gate_evidence/it/capacity_report.csv',[{"service":p['project_name'],"capacity_headroom_pct":it['capacity_headroom_pct'],"continuity_class":it['service_continuity_class'],"support_model":it['support_model']}])
    if it['catalog_status']!='standard':
        d=_doc('Technology Waiver / Convergence Plan'); _kv(d,'Exception',[("Catalog status",it['catalog_status']),("Waiver present",it['waiver_present']),("Expiry days",it['waiver_expiry_days']),("Convergence owner","Enterprise Architecture")]); _bullets(d,'Convergence plan',["Document target standard","Define migration milestone","Remove exception before expiry or re-submit for governance"]); _save(d,case_dir/'gate_evidence/it/technology_waiver.docx')

    # ARCHITECTURE HLD / LLD / flows
    arch=case['architecture_profile']; af=case['architecture']
    hld_svg=case_dir/'gate_evidence/architecture/Architecture_Diagram_Detailed.svg'; hld_png=case_dir/'gate_evidence/architecture/Architecture_Diagram_Detailed.png'; _ensure(hld_svg)
    render_architecture(arch,p,hld_svg,hld_png)
    d=_doc('Low-Level Architecture Notes',f"LLD {v('LLD','lld_version',af['lld_version'])}")
    _kv(d,'Architecture constraints',[("API gateway required",af['api_gateway_required']),("API gateway present",v('LLD','api_gateway_present',af['api_gateway_present'])),("DNS zone",af['dns_zone']),("Latency target ms",af['latency_target_ms']),("Measured latency ms",v('LLD','measured_latency_ms',af['measured_latency_ms'])),("Reversibility",v('LLD','reversibility_status',af['reversibility_status']))])
    _bullets(d,'Design notes',[f"Compute profile: {arch['compute_profile']}",f"Network profile: {arch['network_profile']}",f"Resilience: {arch['resilience_profile']}",f"Data replication: {arch['data_replication']}"])
    _save(d,case_dir/'gate_evidence/architecture/LLD_Architecture_Notes.docx')
    # flows
    flows=[]
    if arch['internet_facing']: flows.append({"source":"Internet Users","destination":arch['edge_services'][0],"protocol":"HTTPS/443","purpose":"User traffic","classification":p['data_classification']})
    if arch['api_management']: flows.append({"source":arch['edge_services'][-1],"destination":"Azure API Management","protocol":"HTTPS/443","purpose":"API policy enforcement","classification":p['data_classification']})
    flows.append({"source":arch['compute_services'][0],"destination":arch['data_services'][0],"protocol":"TLS/Private Link" if arch['private_endpoints'] else "TLS/Public endpoint","purpose":"Application data","classification":p['data_classification']})
    write_csv(case_dir/'gate_evidence/architecture/flow_matrix.csv',flows)
    write_text(case_dir/'gate_evidence/architecture/openapi.yaml',f"openapi: 3.0.3\ninfo:\n  title: {p['project_name']} API\n  version: 1.0.0\npaths:\n  /health:\n    get:\n      responses:\n        '200':\n          description: OK\n  /v1/items:\n    get:\n      security:\n        - oauth2: []\n      responses:\n        '200':\n          description: Success\n")
    write_json(case_dir/'gate_evidence/architecture/data_model.json',{"data_owner":p['business_owner'] if af['data_owner_present'] else None,"classification":p['data_classification'],"personal_data":p['personal_data'],"lineage_complete":af['data_lineage_complete']})
    write_csv(case_dir/'gate_evidence/architecture/data_lineage.csv',[{"source":"Business source","target":arch['data_services'][0],"owner":p['business_owner'] if af['data_owner_present'] else "","lineage_status":"complete" if af['data_lineage_complete'] else "partial"}])
    write_csv(case_dir/'gate_evidence/architecture/ip_plan.csv',[{"cidr":c,"overlap_detected":af['ip_overlap']} for c in af['private_cidrs']])
    write_csv(case_dir/'gate_evidence/architecture/adr_register.csv',[{"adr_id":f"ADR-{i+1:03d}","decision":"Use governed Azure reference pattern","status":"Accepted"} for i in range(af['adr_count'])])
    write_csv(case_dir/'gate_evidence/architecture/architecture_debt.csv',[{"debt_id":f"DEBT-{i+1:03d}","severity":"medium","owner":"Architecture","target":"Post go-live"} for i in range(af['architecture_debt_items'])])

    # SECURITY raw evidence
    sec=case['security']
    resources=[]
    for svc in arch['compute_services']+arch['data_services']+arch.get('platform_services',[])+arch['edge_services']:
        resources.append({"name":svc,"publicNetworkAccess":not arch['private_endpoints'] if svc in arch['data_services'] else arch['internet_facing'],"region":arch['primary_region'],"multiAz":arch['multi_az']})
    write_json(case_dir/'gate_evidence/security/azure_resource_graph.json',{"resources":resources,"private_endpoints":arch['private_endpoints'],"firewall":arch['azure_firewall']})
    write_csv(case_dir/'gate_evidence/security/entra_role_assignments.csv',[{"principal":"workload-mi" if sec['managed_identity'] else "svc-shared","type":"ManagedIdentity" if sec['managed_identity'] else "ServicePrincipal","role":"Contributor" if sec['shared_service_principal'] else "Reader","scope":"subscription" if sec['shared_service_principal'] else "resource-group"}])
    write_json(case_dir/'gate_evidence/security/conditional_access.json',{"mfa":sec['mfa'],"pim":sec['pim'],"policy_state":"enabled" if sec['mfa']=="mandatory" else "reportOnly"})
    write_csv(case_dir/'gate_evidence/security/nsg_rules.csv',[{"priority":100,"direction":"Inbound","source":"Internet" if sec['internet_exposed'] else "CorpNet","destination":"AppSubnet","port":"443","action":"Allow"},{"priority":4096,"direction":"Inbound","source":"*","destination":"*","port":"*","action":"Deny"}])
    write_json(case_dir/'gate_evidence/security/waf_policy.json',{"present":sec['waf_present'],"mode":sec['waf_mode'],"managed_rules":"OWASP","bot_protection":arch.get('semantic_edge',{}).get('bot_protection',False)})
    write_json(case_dir/'gate_evidence/security/sentinel_connectors.json',{"sentinel":sec['sentinel'],"logs_to_siem":sec['logs_to_siem'],"connectors":["Azure Activity","Entra ID","Application"] if sec['logs_to_siem'] else []})
    write_json(case_dir/'gate_evidence/security/firewall_policy.json',{"azure_firewall":arch['azure_firewall'],"default_egress":"deny" if arch['azure_firewall'] else "not_applicable","tls_inspection":arch.get('semantic_network',{}).get('firewall_tls_inspection',False),"rule_collections":["web-egress","private-endpoints","management"] if arch['azure_firewall'] else []})
    write_json(case_dir/'gate_evidence/security/keyvault_configuration.json',{"secrets_in_keyvault":sec['secrets_in_keyvault'],"key_rotation_days":sec['key_rotation_days'],"private_endpoint":arch['private_endpoints'],"managed_identity":sec['managed_identity']})
    write_csv(case_dir/'gate_evidence/security/defender_findings.csv',[{"finding":"Cloud posture assessment","severity":"High" if not arch['private_endpoints'] else "Low","status":"Open" if not arch['private_endpoints'] else "Closed"}])
    vulns=[]
    for i in range(sec['critical_vulns_open']): vulns.append({"id":f"CVE-2026-{8000+i}","severity":"Critical","status":"Open"})
    for i in range(sec['high_vulns_open']): vulns.append({"id":f"CVE-2026-{9000+i}","severity":"High","status":"Open"})
    if not vulns: vulns=[{"id":"CVE-2026-0001","severity":"Medium","status":"Remediated"}]
    write_csv(case_dir/'gate_evidence/security/vulnerability_scan.csv',vulns)
    d=_doc('Penetration Test Report'); _kv(d,'Summary',[("Status",v('PENTEST','status',sec['pentest_status'])),("Age days",v('PENTEST','age',sec['pentest_age_days'])),("Critical findings",sec['critical_vulns_open']),('High findings',sec['high_vulns_open'])]); _bullets(d,'Tester comments',["External and authenticated test paths executed where applicable.","Findings require independent verification against the vulnerability export."]); _save(d,case_dir/'gate_evidence/security/pentest_report.docx')
    d=_doc('Threat Model'); _bullets(d,'Trust boundaries',["Internet / edge boundary","Identity provider boundary","Application / data boundary","Primary / DR region boundary" if arch['multi_region'] else "Single-region resilience boundary"]); _bullets(d,'Threat scenarios',["Credential misuse","Public endpoint exposure","Data exfiltration","Supply-chain compromise","Availability failure"]); _save(d,case_dir/'gate_evidence/security/threat_model.docx')

    # TECH READINESS
    tr=case['tech_readiness']
    write_csv(case_dir/'gate_evidence/tech_readiness/backup_jobs.csv',[{"job":"daily-backup","enabled":tr['backup_enabled'],"last_status":"Succeeded" if tr['backup_enabled'] else "NotConfigured","retention_days":35}])
    write_csv(case_dir/'gate_evidence/tech_readiness/restore_test.csv',[{"tested":tr['restore_tested'],"target_rto_hours":tr['target_rto_hours'],"measured_restore_hours":tr['measured_restore_hours'],"target_rpo_minutes":tr['target_rpo_minutes'],"measured_data_loss_minutes":tr['measured_data_loss_minutes']}])
    write_csv(case_dir/'gate_evidence/tech_readiness/failover_test.csv',[{"dr_required":tr['dr_required'],"tested":tr['dr_tested'],"secondary_region":arch['secondary_region'] or "N/A","result":"PASS" if (not tr['dr_required'] or tr['dr_tested']) else "NOT_RUN"}])
    write_csv(case_dir/'gate_evidence/tech_readiness/load_test.csv',[{"peak_test_pct":tr['load_test_pct_of_peak'],"p95_ms":tr['p95_ms'],"test_completion_pct":tr['test_completion_pct']}])
    write_json(case_dir/'gate_evidence/tech_readiness/ci_pipeline.json',{"unit_test_coverage_pct":tr['unit_test_coverage_pct'],"critical_static_findings":tr['critical_static_findings'],"build_reproducible":True})
    write_csv(case_dir/'gate_evidence/tech_readiness/rollback_test.csv',[{"rollback_tested":tr['rollback_tested'],"result":"PASS" if tr['rollback_tested'] else "NOT_RUN"}])
    write_csv(case_dir/'gate_evidence/tech_readiness/slo_sli.csv',[{"sli":"availability","target":"99.9%","observed":"99.95%" if tr['open_blockers']==0 else "99.6%"},{"sli":"p95_latency_ms","target":case['architecture']['latency_target_ms'],"observed":tr['p95_ms']}])
    write_csv(case_dir/'gate_evidence/tech_readiness/monitor_alerts.csv',[{"alert":"Service availability","configured":tr['observability_metrics'],"severity":"critical"},{"alert":"Backup failure","configured":tr['backup_enabled'],"severity":"high"},{"alert":"Security log ingestion","configured":case['security']['logs_to_siem'],"severity":"high"}])
    d=_doc('Production Runbook & Handover'); _kv(d,'Operations',[("Runbook status",v('RUNBOOK','status',tr['runbook_status'])),("On-call defined",v('RUNBOOK','on_call',tr['on_call_defined'])),("Handover signed",v('RUNBOOK','handover',tr['handover_signed'])),("Service owner",p['service_owner'])]); _bullets(d,'Operational procedures',["Incident triage and escalation","Backup verification","Restore execution","Rollback deployment","Failover to secondary region" if arch['multi_region'] else "Same-region recovery"]); _save(d,case_dir/'gate_evidence/tech_readiness/production_runbook.docx')

    # PROCUREMENT
    pr=case['procurement']
    d=_doc('Request for Proposal',f"RFP for {p['project_name']}"); _kv(d,'Tender',[("Requirements",pr['rfp_requirements']),("Competition",pr['competition']),("Budget",pr['purchasing_budget_eur'])]); _bullets(d,'Mandatory criteria',["Security evidence package","Data location disclosure","Support SLA","Exit and data export","Sanctions clearance"]); _save(d,case_dir/'gate_evidence/procurement/RFP.docx')
    write_csv(case_dir/'gate_evidence/procurement/vendor_offers.csv',pr['offers'])
    write_csv(case_dir/'gate_evidence/procurement/scoring_matrix.csv',[{"vendor":o['vendor'],"functional":o['functional_score'],"security":o['security_score'],"tco_3y_eur":o['tco_3y_eur'],"weighted_score":o['weighted_score'],"selected":o['vendor']==pr['selected_vendor']} for o in pr['offers']])
    write_csv(case_dir/'gate_evidence/procurement/due_diligence.csv',[{"vendor":o['vendor'],"mandatory_failed":o['mandatory_criteria_failed'],"sanctions":o['sanctions'],"due_diligence":o['due_diligence'],"references_checked":o['references_checked']} for o in pr['offers']])
    write_csv(case_dir/'gate_evidence/procurement/pricing_tco.csv',[{"vendor":o['vendor'],"tco_3y_eur":o['tco_3y_eur'],"annualized_eur":round(o['tco_3y_eur']/3,2),"selected":o['vendor']==pr['selected_vendor']} for o in pr['offers']])
    d=_doc('Supplier Evidence Request'); _kv(d,'Request',[("Selected vendor",pr['selected_vendor']),("Project",p['project_name']),("Requested by","Procurement / Security")]); _bullets(d,'Evidence requested',["Current ISO/SOC reports or equivalent","Subprocessor inventory","Data-location statement","Security incident notification terms","BCP/DR evidence","Exit and deletion process"]); _save(d,case_dir/'gate_evidence/procurement/vendor_evidence_request.docx')

    # LEGAL
    le=case['legal']
    d=_doc('Master Services Agreement',f"Synthetic contract {le['contract_version']}")
    clauses=[
        ('1. Services',f"Supplier shall provide the services described in the Order Form for {p['project_name']}."),
        ('2. Term','Initial term 36 months with renewal subject to the Order Form.'),
        ('3. Service Levels',f"Service credits and penalties: {'included' if le['sla_penalties'] else 'not included'}."),
        ('4. Data Processing',f"DPA status: {v('MSA','dpa_status',le['dpa_status'])}. Transfer clause: {le['data_transfer_clause']}."),
        ('5. Confidentiality','Each party protects confidential information using reasonable safeguards.'),
        ('6. Security',f"Security logs export: {v('MSA','log_export_clause',le['log_export_clause'])}."),
        ('7. Audit Rights',f"Customer audit rights: {v('MSA','audit_right',le['audit_right'])}."),
        ('8. Intellectual Property',f"IP ownership: {v('MSA','ip_ownership',le['ip_ownership'])}."),
        ('9. Exit and Reversibility',f"Exit assistance: {v('MSA','exit_days',le['exit_assistance_days'])} days."),
        ('10. Insurance',f"Cyber insurance: EUR {le['cyber_insurance_m']}m. Certificate valid: {le['insurance_valid']}."),
        ('11. Liability',f"Aggregate liability cap: {le['liability_cap_multiplier']}x annual fees. Security carve-out: {le['security_carveout']}."),
        ('12. Governing Law',le['governing_law']),
    ]
    for title,body in clauses: d.add_heading(title,level=1); d.add_paragraph(body); d.add_paragraph('Supplier comment: Clause accepted subject to final commercial approval.' if 'Liability' not in title else 'Supplier comment: Current liability position is commercially sensitive.')
    _save(d,case_dir/'gate_evidence/legal/MSA.docx')
    d=_doc('Data Processing Addendum'); _kv(d,'Data processing',[("Status",v('DPA','status',le['dpa_status'])),("Roles","Customer controller / Supplier processor"),("Primary region",p['primary_region']),("Personal data",p['personal_data'])]); _save(d,case_dir/'gate_evidence/legal/DPA.docx')
    d=_doc('SLA Annex'); _kv(d,'Service levels',[("Availability","99.9%"),("P1 response","30 minutes"),("Penalties",le['sla_penalties']),("Recovery target",f"RTO {arch['rto_hours']}h / RPO {arch['rpo_minutes']}m")]); _save(d,case_dir/'gate_evidence/legal/SLA_Annex.docx')
    d=_doc('Insurance Certificate'); _kv(d,'Certificate',[("Valid",v('INSURANCE','valid',le['insurance_valid'])),("Cyber coverage EUR m",le['cyber_insurance_m']),("Insured","Selected Supplier")]); _save(d,case_dir/'gate_evidence/legal/Insurance_Certificate.docx')
    write_csv(case_dir/'gate_evidence/legal/signing_authority.csv',[{"signatory":p['sponsor'],"authority_valid":le['signing_authority_valid'],"limit_eur":p['budget_requested_eur']}])
    write_json(case_dir/'gate_evidence/legal/contract_playbook.json',{"required":{"DPA":"signed if personal data","audit_right":"full or limited","exit_assistance_days":">=30","security_log_export":"included for high/critical services","signing_authority":"valid"}})
    d=_doc('Contract Negotiation Log'); _bullets(d,'Open negotiation points',[f"Liability cap currently {le['liability_cap_multiplier']}x annual fees.",f"Audit right currently: {le['audit_right']}.",f"Exit assistance: {le['exit_assistance_days']} days.",f"Security log export: {le['log_export_clause']}.",f"DPA status: {le['dpa_status']}."]); _save(d,case_dir/'gate_evidence/legal/negotiation_log.docx')

    # COMPLIANCE
    co=case['compliance']
    write_csv(case_dir/'gate_evidence/compliance/data_inventory.csv',[{"dataset":"Core business data","classification":p['data_classification'],"personal_data":p['personal_data'],"actual_region":p['primary_region'],"required_residency":co['required_residency']}])
    write_csv(case_dir/'gate_evidence/compliance/ropa.csv',[{"processing":"Core service","purpose":"Business operations","controller":p['business_unit'],"processor":pr['selected_vendor'],"personal_data":p['personal_data'],"retention_years":co['retention_years']}])
    d=_doc('Data Protection Impact Assessment'); _kv(d,'DPIA',[("Required",co['dpia_required']),("Status",v('DPIA','status',co['dpia_status'])),("Personal data",p['personal_data']),("AI enabled",p['ai_enabled']),("Residual risk owner",p['risk_owner'])]); _save(d,case_dir/'gate_evidence/compliance/DPIA.docx')
    write_csv(case_dir/'gate_evidence/compliance/regulatory_mapping.csv',[{"regulation":r,"mapping_status":co['regulatory_mapping'],"owner":"GRC"} for r in co['applicable_regulations']])
    write_csv(case_dir/'gate_evidence/compliance/retention_schedule.csv',[{"record":"Application data","retention_years":co['retention_years'],"legal_hold":co['legal_hold_supported'],"deletion_enforced":co['deletion_enforced']}])
    write_csv(case_dir/'gate_evidence/compliance/accessibility_report.csv',[{"scope":"Web / user interface","status":co['accessibility_status'],"standard":"WCAG 2.2 AA"}])
    write_csv(case_dir/'gate_evidence/compliance/control_matrix.csv',[{"obligation":r,"control_mapping_status":co['regulatory_mapping'],"evidence_owner":"GRC","evidence":"See regulatory_mapping.csv"} for r in co['applicable_regulations']])
    write_csv(case_dir/'gate_evidence/compliance/export_control.csv',[{"status":co['export_control'],"review_owner":"Compliance","project_region":p['primary_region']}])
    d=_doc('AI Impact & Ethics Assessment'); _kv(d,'Assessment',[("AI enabled",p['ai_enabled']),("Ethics review",co['ethics_review']),("High-impact decision",False),("Human override available",True)]); _save(d,case_dir/'gate_evidence/compliance/AI_Impact_Assessment.docx')

    # Shared evidence provenance registry.
    provenance=[]
    for n in graph['nodes']:
        provenance.append({"evidence_id":n['evidence_id'],"version":n['version'],"age_days":n['age_days'],"authoritative":n['authoritative'],"status":n['public_status'],"path":n['path']})
    write_csv(case_dir/'shared/evidence_provenance.csv',provenance)

    # Shared action register generated from reference outcomes, but without hidden judgement rationale.
    rows=[]
    for rr in reference_results:
        for a in rr['required_actions']:
            rows.append({"occurrence_id":rr['occurrence_id'],"gate":rr['gate'],"finding_id":a['finding_id'],"action":a['action'],"owner":a['owner'],"status":"OPEN","due_date":p['target_go_live']})
    if not rows: rows=[{"occurrence_id":"N/A","gate":"N/A","finding_id":"NONE","action":"No open actions","owner":p['project_manager'],"status":"CLOSED","due_date":p['target_go_live']}]
    write_csv(case_dir/'shared/action_register.csv',rows)

    make_review_requests(case_dir,case,graph,occurrences)
