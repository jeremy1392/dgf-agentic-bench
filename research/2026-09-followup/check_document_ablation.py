"""Offline counterexample to scoring a DOCX-only ablation with the frozen reference.

Generate two dossier variants that differ only in the selected vendor's due
diligence status. Compare the extracted paragraphs and tables of every DOCX.
Original datasets and model results are never changed; no inference calls.
"""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import sys

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[1]
SOURCE=ROOT/'experiments/reproduction_check_20260923/verified_inputs/benchmark_source'
sys.path.insert(0,str(SOURCE))
from document_factory import emit_all
from evaluator import evaluate_gate
from docx import Document


def text(path):
    d=Document(path)
    return '\n'.join([p.text for p in d.paragraphs]+[' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows])


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--dataset',type=Path,default=ROOT/'experiments/preflight_balanced_300_20260922/dataset')
    ap.add_argument('--output-dir',type=Path,default=ROOT/'experiments/document_ablation_preflight_20260924')
    ap.add_argument('--report',type=Path,default=HERE/'document_ablation_preflight.json');args=ap.parse_args()
    for case in sorted(args.dataset.glob('DGF-BUY-*')):
        truth=json.loads((case/'99_hidden_ground_truth.json').read_text(encoding='utf-8'))
        occurrence=next(x for x in json.loads((case/'01_route_manifest.json').read_text())['occurrences'] if x['gate']=='procurement')
        original=truth['canonical_truth'];ref=evaluate_gate(original,'procurement',occurrence['phase'])
        if ref['disposition']=='GO':break
    else:raise ValueError('No GO procurement case found')
    variants={}
    for status in ['complete','pending']:
        canonical=copy.deepcopy(original)
        selected=next(x for x in canonical['procurement']['offers'] if x['vendor']==canonical['procurement']['selected_vendor'])
        selected['due_diligence']=status
        folder=args.output_dir/status
        emit_all(folder,canonical,truth['evidence_graph'],[occurrence],[])
        variants[status]={'reference':evaluate_gate(canonical,'procurement',occurrence['phase']),
                         'docs':{p.relative_to(folder).as_posix():text(p) for p in sorted(folder.rglob('*.docx'))}}
    assert variants['complete']['docs']==variants['pending']['docs']
    assert variants['complete']['reference']['disposition']=='GO'
    assert variants['pending']['reference']['disposition']=='REWORK'
    assert [x['id'] for x in variants['pending']['reference']['findings']]==['PROC-DD-001']
    complete_dir=args.output_dir/'complete';pending_dir=args.output_dir/'pending'
    changed_records=sorted(p.relative_to(complete_dir).as_posix() for p in complete_dir.rglob('*')
        if p.is_file() and p.suffix!='.docx'
        and p.read_bytes()!=(pending_dir/p.relative_to(complete_dir)).read_bytes())
    assert changed_records==['gate_evidence/procurement/due_diligence.csv',
        'gate_evidence/procurement/review_facts.json','gate_evidence/procurement/vendor_offers.csv']
    result={'status':'DOCX_ONLY_ABLATION_NOT_IDENTIFIABLE_UNDER_FROZEN_REFERENCE',
        'original_case':case.name,'phase':occurrence['phase'],'counterfactual_field':'procurement.offers[selected_vendor].due_diligence',
        'variants':{k:{'value':k,'reference_disposition':v['reference']['disposition'],
                       'finding_ids':[x['id'] for x in v['reference']['findings']]} for k,v in variants.items()},
        'identical_docx_text_count':len(variants['complete']['docs']),
        'document_text_sha256':{k:hashlib.sha256(v.encode()).hexdigest() for k,v in variants['complete']['docs'].items()},
        'differing_non_docx_evidence':changed_records,
        'interpretation':'The same DOCX paragraphs and tables correspond to two different required procurement decisions. A Word-only reviewer cannot identify the true due diligence state from these documents. This counterexample does not establish insufficiency of a broader document-plus-CSV condition.',
        'additional_contract_blocker':'Frozen evidence scoring requires findings to cite REVIEW_FACTS_GATE. Removing that evidence while retaining the strict score makes citation conformity impossible even when a decision can be inferred. A common versioned evidence contract must be designed for every ablation cell.',
        'model_calls':0,'canonical_truth_use':'Used offline to construct and evaluate the counterfactual, never supplied to a model.'}
    args.report.write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8');print(json.dumps(result,indent=2))


if __name__=='__main__':main()
