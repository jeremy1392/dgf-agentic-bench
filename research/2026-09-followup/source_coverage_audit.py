"""Conservative offline certification of non-REVIEW_FACTS source coverage.

Read-only audit of all 300 dossiers. Explicit, template-specific decoders are
matched to the frozen evaluator's AST fields. Failure to certify is never called
proof of absence. No model calls, rescoring, generator changes or dataset edits.
"""
import argparse
import ast
import collections
import copy
import csv
import hashlib
import json
from pathlib import Path
import re
import sys
from unittest.mock import patch

from docx import Document

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SOURCE = ROOT/'experiments/reproduction_check_20260923/verified_inputs/benchmark_source'
sys.path.insert(0, str(SOURCE))
from benchmark_protocol import GATES, source_fingerprint


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8'))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inventory():
    """Resolve the frozen simple local aliases and procurement selector via AST."""
    tree = ast.parse((SOURCE/'evaluator.py').read_text(encoding='utf-8'))
    output = {}
    for fn in [x for x in tree.body if isinstance(x, ast.FunctionDef) and x.name in
               {'evaluate_'+('tech' if g == 'tech_readiness' else g) for g in GATES}]:
        gate = fn.name.removeprefix('evaluate_')
        if gate == 'tech': gate = 'tech_readiness'
        aliases = {'case': ()}
        def resolve(node):
            if isinstance(node, ast.Name): return aliases.get(node.id)
            if isinstance(node, ast.Subscript) and isinstance(node.slice, ast.Constant) and isinstance(node.slice.value, str):
                parent = resolve(node.value)
                if parent is not None: return (*parent, node.slice.value)
            return None
        for node in fn.body:
            if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
                value = resolve(node.value)
                if value is not None: aliases[node.targets[0].id] = value
        # next(x for x in f['offers'] if x['vendor']==f['selected_vendor']).
        for node in ast.walk(fn):
            if isinstance(node, ast.GeneratorExp):
                for gen in node.generators:
                    path = resolve(gen.iter)
                    if path and isinstance(gen.target, ast.Name): aliases[gen.target.id] = (*path, '[*]')
        for node in fn.body:
            if isinstance(node, ast.Assign) and isinstance(node.value, ast.Call) and isinstance(node.value.func, ast.Name) and node.value.func.id == 'next':
                gen = node.value.args[0]
                if isinstance(gen, ast.GeneratorExp):
                    path = resolve(gen.elt)
                    if path and isinstance(node.targets[0], ast.Name): aliases[node.targets[0].id] = (*path[:-1], '[selected_vendor]')
        paths = collections.defaultdict(set)
        for node in ast.walk(fn):
            if isinstance(node, ast.Subscript):
                path = resolve(node)
                if path and len(path) >= 2: paths['.'.join(path)].add(node.lineno)
        # Containers are not additional scalar requirements.
        paths = {p: lines for p, lines in paths.items() if not any(q.startswith(p+'.') for q in paths)}
        output[gate] = [{'field': p, 'evaluator_lines': sorted(lines),
                        'role': 'output_metadata' if p == 'project.risk_owner' else
                                'supplier_selector' if p in ('procurement.selected_vendor', 'procurement.offers.[*].vendor') else 'rule_input'}
                       for p, lines in sorted(paths.items())]
    assert set(output) == set(GATES)
    return output


def spec(eid, kind, key, cast='str', **kwargs):
    return dict(evidence_id=eid, kind=kind, key=key, cast=cast, **kwargs)


def decoders():
    d = {}
    def add(field, *sources): d[field] = list(sources)
    add('project.risk_owner', spec('PROJECT_CHARTER', 'kv', 'Risk owner'))
    add('project.business_criticality', spec('PROJECT_CHARTER', 'kv', 'Criticality'))
    add('project.data_classification', spec('PROJECT_CHARTER', 'kv', 'Data classification'), spec('DATA_INVENTORY', 'csv', 'classification'))
    add('project.personal_data', spec('DATA_INVENTORY', 'csv', 'personal_data', 'bool'), spec('DPA', 'kv', 'Personal data', 'bool'), spec('DATA_MODEL', 'json', 'personal_data', 'bool'))
    for field, eid, col, cast in [
        ('catalog_status','TECH_CATALOG','catalog_status','str'), ('waiver_present','TECH_CATALOG','waiver_present','bool'),
        ('license_compliant','LICENSE_POSITION','license_compliant','bool'), ('cmdb_record_present','CMDB_EXPORT','record_present','bool'),
        ('run_owner_present','CMDB_EXPORT','run_owner','nonempty'), ('technology_eol_months','LIFECYCLE_REGISTER','eol_months','int'),
        ('capacity_headroom_pct','CAPACITY_REPORT','capacity_headroom_pct','int'), ('change_record_status','ITSM_CHANGE','status','str')]:
        add('it.'+field, spec(eid, 'json' if eid=='TECH_CATALOG' else 'csv', col, cast))
    for field, label, cast in [('api_gateway_required','API gateway required','bool'),('api_gateway_present','API gateway present','bool'),
        ('latency_target_ms','Latency target ms','int'),('measured_latency_ms','Measured latency ms','int'),('reversibility_status','Reversibility','str')]:
        add('architecture.'+field, spec('LLD','kv',label,cast))
    add('architecture.ip_overlap', spec('IP_PLAN','csv','overlap_detected','bool'))
    add('architecture.data_owner_present',spec('DATA_MODEL','json','data_owner','nonempty'),spec('DATA_LINEAGE','csv','owner','nonempty'))
    add('architecture_profile.multi_az',spec('AZURE_RESOURCE_GRAPH','resources_consistent','multiAz','bool'))
    add('architecture_profile.multi_region',spec('THREAT_MODEL','exact_choice','Primary / DR region boundary','bool',false_text='Single-region resilience boundary'))
    for field, eid, col, cast in [('waf_present','WAF_POLICY','present','bool'),('waf_mode','WAF_POLICY','mode','str'),
        ('private_endpoints','AZURE_RESOURCE_GRAPH','private_endpoints','bool'),('mfa','CONDITIONAL_ACCESS','mfa','str'),
        ('logs_to_siem','SENTINEL_STATUS','logs_to_siem','bool'),('key_rotation_days','KEYVAULT_CONFIG','key_rotation_days','nullable_int')]:
        add('security.'+field,spec(eid,'json',col,cast))
    add('security.internet_exposed',spec('NSG_RULES','csv_where','source','internet_source',where_key='priority',where_value='100'))
    add('security.shared_service_principal',spec('IAM_EXPORT','identity_encoding','role','bool'))
    add('security.critical_vulns_open',spec('VULN_SCAN','count_critical_open','severity','int'),spec('PENTEST','kv','Critical findings','int'))
    add('security.pentest_status',spec('PENTEST','kv','Status'))
    for field, eid, col, cast in [('backup_enabled','BACKUP_JOBS','enabled','bool'),('restore_tested','RESTORE_TEST','tested','bool'),
        ('target_rto_hours','RESTORE_TEST','target_rto_hours','float'),('measured_restore_hours','RESTORE_TEST','measured_restore_hours','float'),
        ('target_rpo_minutes','RESTORE_TEST','target_rpo_minutes','float'),('measured_data_loss_minutes','RESTORE_TEST','measured_data_loss_minutes','float'),
        ('dr_required','FAILOVER_TEST','dr_required','bool'),('dr_tested','FAILOVER_TEST','tested','bool'),
        ('load_test_pct_of_peak','LOAD_TEST','peak_test_pct','int'),('critical_static_findings','CI_PIPELINE','critical_static_findings','int'),
        ('rollback_tested','ROLLBACK_TEST','rollback_tested','bool')]:
        add('tech_readiness.'+field,spec(eid,'json' if eid=='CI_PIPELINE' else 'csv',col,cast))
    for field, label, cast in [('runbook_status','Runbook status','str'),('on_call_defined','On-call defined','bool'),('handover_signed','Handover signed','bool')]:
        add('tech_readiness.'+field,spec('RUNBOOK','kv',label,cast))
    add('procurement.purchasing_budget_eur',spec('RFP','kv','Budget','int'))
    add('procurement.selected_vendor',spec('SCORING_MATRIX','selected_name','vendor'),spec('PRICING_TCO','selected_name','vendor'),spec('VENDOR_EVIDENCE_REQUEST','kv','Selected vendor'))
    add('procurement.offers.[*].vendor',spec('VENDOR_OFFERS','vendor_names','vendor','identity'))
    for field, cast in [('mandatory_criteria_failed','int'),('sanctions','str'),('due_diligence','str'),('tco_3y_eur','int'),('references_checked','bool')]:
        eid='PRICING_TCO' if field=='tco_3y_eur' else 'DUE_DILIGENCE'
        col='mandatory_failed' if field=='mandatory_criteria_failed' else field
        add('procurement.offers.[selected_vendor].'+field,spec('VENDOR_OFFERS','selected_row',field,cast),spec(eid,'selected_row',col,cast))
    add('legal.dpa_status',spec('DPA','kv','Status'),spec('MSA','regex',r'^DPA status: ([^.]+)\. Transfer clause:'),spec('NEGOTIATION_LOG','regex',r'^DPA status: ([^.]+)\.$'))
    for field, regex, cast in [
        ('liability_cap_multiplier',r'^Aggregate liability cap: ([0-9.]+|unlimited)x annual fees\.', 'liability'),
        ('security_carveout',r'^Aggregate liability cap: (?:[0-9.]+|unlimited)x annual fees\. Security carve-out: (True|False)\.$','bool'),
        ('audit_right',r'^Customer audit rights: ([^.]+)\.$','str'), ('exit_assistance_days',r'^Exit assistance: ([0-9]+) days\.$','int'),
        ('ip_ownership',r'^IP ownership: ([^.]+)\.$','str'), ('insurance_valid',r'^Cyber insurance: EUR [0-9.]+m\. Certificate valid: (True|False)\.$','bool'),
        ('log_export_clause',r'^Security logs export: ([^.]+)\.$','str')]:
        add('legal.'+field,spec('MSA','regex',regex,cast))
    for field, regex, cast in [('liability_cap_multiplier',r'^Liability cap currently ([0-9.]+|unlimited)x annual fees\.$','liability'),
        ('audit_right',r'^Audit right currently: ([^.]+)\.$','str'), ('exit_assistance_days',r'^Exit assistance: ([0-9]+) days\.$','int'),
        ('log_export_clause',r'^Security log export: ([^.]+)\.$','str')]:
        d['legal.'+field].append(spec('NEGOTIATION_LOG','regex',regex,cast))
    d['legal.insurance_valid'].append(spec('INSURANCE','kv','Valid','bool'))
    add('legal.signing_authority_valid',spec('SIGNING_AUTHORITY','csv','authority_valid','bool'))
    for field,eid,key,cast in [('dpia_required','DPIA','Required','bool'),('dpia_status','DPIA','Status','str'),
        ('regulatory_mapping','REG_MAPPING','mapping_status','str'),('accessibility_status','ACCESSIBILITY','status','str'),('export_control','EXPORT_CONTROL','status','str')]:
        add('compliance.'+field,spec(eid,'kv' if eid=='DPIA' else 'csv',key,cast))
    add('general.strategic_alignment',spec('PORTFOLIO_SNAPSHOT','csv','strategic_alignment'),spec('BUSINESS_CASE','csv_metric','Strategic_alignment'))
    for field,col in [('budget_requested_eur','requested_eur'),('budget_approved_eur','approved_eur')]:
        add('general.'+field,spec('BUDGET_APPROVAL','csv',col,'int'),spec('PORTFOLIO_SNAPSHOT','csv',col,'int'))
    add('general.roi',spec('BUSINESS_CASE','csv_metric','ROI','float'),spec('COMMITTEE_BRIEFING','kv','ROI','float'))
    add('general.change_plan_status',spec('BENEFITS_PLAN','kv','Change plan'))
    return d


def convert(value, kind):
    if kind == 'identity': return value
    if kind == 'liability':
        if value == 'unlimited': return value
        return convert(value, 'float')
    if kind == 'nonempty': return value not in ('', None)
    if kind == 'bool':
        if type(value) is bool: return value
        if value in ('True','False'): return value == 'True'
        raise ValueError('Not an explicit boolean')
    if kind == 'internet_source':
        if value not in ('Internet','CorpNet'): raise ValueError('Unexpected NSG source encoding')
        return value == 'Internet'
    if kind == 'nullable_int' and value is None: return None
    if kind in ('int','nullable_int'):
        if isinstance(value,bool) or float(value) != int(float(value)): raise ValueError('Not an exact integer')
        return int(float(value))
    if kind == 'float':
        if isinstance(value,bool): raise ValueError('Boolean is not a number')
        return float(value)
    if kind == 'str' and isinstance(value,str): return value
    raise ValueError('Unexpected decoded type')


def equal(a,b):
    if isinstance(a,bool) or isinstance(b,bool): return type(a) is type(b) and a == b
    return a == b


class Sources:
    def __init__(self, folder):
        self.folder = folder
        self.nodes = {n['evidence_id']:n for n in read_json(folder/'02_evidence_graph.json')['nodes'] if not n['evidence_id'].startswith('REVIEW_FACTS_')}
        self.visibility = read_json(folder/'05_phase_visibility.json')
        self.cache = {}
        self.hashes = {}

    def load(self,eid):
        if eid in self.cache: return self.cache[eid]
        node = self.nodes[eid]; path = self.folder/node['path']
        assert self.folder.resolve() in path.resolve().parents
        self.hashes[node['path']] = digest(path)
        if path.suffix == '.json': content = read_json(path)
        elif path.suffix == '.csv':
            with path.open(encoding='utf-8-sig',newline='') as f: content = list(csv.DictReader(f))
            if len(content)>200: raise ValueError('Reader row truncation would need separate certification')
        elif path.suffix == '.docx':
            doc=Document(path)
            content={'paragraphs':[p.text.strip() for p in doc.paragraphs if p.text.strip()],
                     'rows':[[c.text.strip() for c in row.cells] for table in doc.tables for row in table.rows]}
            text='\n'.join(content['paragraphs']+[' | '.join(r) for r in content['rows']])
            if len(text)>40000: raise ValueError('Reader character truncation would need separate certification')
        else: raise ValueError('Unsupported source format')
        self.cache[eid]=content
        return content

    def selected(self):
        rows=self.load('SCORING_MATRIX')
        vals=[row['vendor'] for row in rows if convert(row['selected'],'bool')]
        if len(vals)!=1: raise ValueError('Public selection is not unique')
        return vals[0]

    def decode(self,s):
        content=self.load(s['evidence_id']); key=s['key']; kind=s['kind']; dependencies=[s['evidence_id']]
        if kind=='json': values=[content[key]]
        elif kind=='resources_consistent': values=[r[key] for r in content['resources']]
        elif kind=='csv': values=[r[key] for r in content]
        elif kind=='csv_where': values=[r[key] for r in content if r[s['where_key']]==s['where_value']]
        elif kind=='csv_metric': values=[r['value'] for r in content if r['metric']==key]
        elif kind=='selected_name': values=[r[key] for r in content if convert(r['selected'],'bool')]
        elif kind=='vendor_names': values=[sorted(r[key] for r in content)]
        elif kind=='selected_row':
            selected=self.selected(); dependencies.append('SCORING_MATRIX')
            values=[r[key] for r in content if r['vendor']==selected]
            if len(values)!=1: raise ValueError('Selected vendor does not identify one source row')
        elif kind=='kv': values=[r[1] for r in content['rows'] if len(r)==2 and r[0]==key]
        elif kind=='regex': values=[m.group(1) for p in content['paragraphs'] if (m:=re.search(key,p))]
        elif kind=='exact_choice':
            yes=key in content['paragraphs']; no=s['false_text'] in content['paragraphs']
            if yes==no: raise ValueError('Template choice is missing or ambiguous')
            values=[yes]
        elif kind=='count_critical_open': values=[sum(r['severity']=='Critical' and r['status']=='Open' for r in content)]
        elif kind=='identity_encoding':
            values=[]
            for row in content:
                if (row['role'],row['scope'])==('Contributor','subscription'): values.append(True)
                elif (row['role'],row['scope'])==('Reader','resource-group'): values.append(False)
                else: raise ValueError('Unexpected frozen identity template')
        else: raise ValueError('No decoder')
        values=[convert(v,s['cast']) for v in values]
        if not values or not all(equal(values[0],v) for v in values): raise ValueError('Missing or inconsistent labelled values')
        return values[0], dependencies

    def availability(self,eid,gate,phase):
        node=self.nodes[eid]
        return {'evidence_id':eid,'in_gate_scope':gate in node['consumers'],
                'in_phase':eid in self.visibility.get(phase,[]),'public_status':node['public_status'],
                'authoritative':node['authoritative'],
                'direct_readable':gate in node['consumers'] and eid in self.visibility.get(phase,[]) and node['public_status']=='AVAILABLE'}


def expected_value(snapshot, field):
    path=field.split('.')
    if '[selected_vendor]' in path:
        offer=next(x for x in snapshot['procurement']['offers'] if x['vendor']==snapshot['procurement']['selected_vendor'])
        return offer[path[-1]]
    if '[*]' in path: return sorted(x['vendor'] for x in snapshot['procurement']['offers'])
    value=snapshot
    for key in path: value=value[key]
    return value


def omission_proofs(dataset):
    """Capture exact generator payloads; no evidence files are written.

    Only REVIEW_FACTS changes under each intervention. Document paragraphs and
    tables, all other serialized files, and diagram renderer inputs stay equal.
    """
    import document_factory as factory
    from evaluator import evaluate_gate
    reports=[]
    for gate,field,values in [('it','duplicate_capability',(False,True)),('compliance','audit_trail',('complete','missing'))]:
        for folder in sorted(dataset.glob('DGF-*')):
            hidden=read_json(folder/'99_hidden_ground_truth.json')
            route=read_json(folder/'01_route_manifest.json')['occurrences']
            occ=next((o for o in route if o['gate']==gate),None)
            if occ and evaluate_gate(hidden['canonical_truth'],gate,occ['phase'])['disposition']=='GO': break
        else: raise ValueError('No clean counterexample base found')
        variants=[]
        for value in values:
            case=copy.deepcopy(hidden['canonical_truth']);case[gate][field]=value
            payloads={}
            def capture(path,value): payloads[Path(path).relative_to(HERE).as_posix()]=copy.deepcopy(value)
            def save_doc(doc,path):
                capture(path,{'paragraphs':[p.text for p in doc.paragraphs],
                              'tables':[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]})
            def render(arch,project,svg,png):
                capture(svg,{'renderer_input_architecture':arch,'renderer_input_project':project})
                capture(png,{'renderer_input_architecture':arch,'renderer_input_project':project})
            with patch.multiple(factory,write_json=capture,write_text=capture,
                                write_csv=lambda path,rows,fieldnames=None:capture(path,{'rows':rows,'fieldnames':fieldnames}),
                                _save=save_doc,_ensure=lambda path:path,render_architecture=render):
                factory.emit_all(HERE,case,hidden['evidence_graph'],route,[])
            variants.append({'value':value,'reference':evaluate_gate(case,gate,occ['phase']),'payloads':payloads})
        a,b=variants
        assert a['payloads'].keys()==b['payloads'].keys()
        changed=[p for p in a['payloads'] if a['payloads'][p]!=b['payloads'][p]]
        assert changed==[f'gate_evidence/{gate}/review_facts.json']
        assert a['reference']['disposition']=='GO' and b['reference']['disposition']=='REWORK'
        reports.append({'field':gate+'.'+field,'base_case':folder.name,'phase':occ['phase'],
                        'variants':[{'value':x['value'],'decision':x['reference']['disposition'],
                                     'findings':[f['id'] for f in x['reference']['findings']]} for x in variants],
                        'changed_payloads':changed,'identical_non_snapshot_payloads':sum('review_facts.json' not in p for p in a['payloads']),
                        'comparison':'All emitted JSON/CSV/text data and DOCX paragraphs/tables; identical renderer inputs for SVG/PNG. DOCX zip metadata and actual renderer output bytes are not compared.',
                        'hidden_truth_use':'Offline counterfactual construction only; not a source decoder or model input.'})
    return reports


def annotate_tool_boundary(output):
    """Record the distinct synthetic-tool interface; it is not part of coverage."""
    path=SOURCE/'synthetic_environment.py'
    tree=ast.parse(path.read_text(encoding='utf-8'))
    methods={n.name:n for n in ast.walk(tree) if isinstance(n,ast.FunctionDef)}
    tools={}
    for name in ('get_cmdb_record','get_regulatory_applicability','request_evidence'):
        fn=methods[name]
        keys=sorted({key.value for node in ast.walk(fn) if isinstance(node,ast.Return) and isinstance(node.value,ast.Dict)
                     for key in node.value.keys if isinstance(key,ast.Constant) and isinstance(key.value,str)})
        tools[name]={'source_line':fn.lineno,'returned_dict_keys':keys,
                     'calls_factual_snapshot':any(isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id=='factual_snapshot' for n in ast.walk(fn))}
    assert 'duplicate_capability' not in tools['get_cmdb_record']['returned_dict_keys']
    assert 'audit_trail' not in tools['get_regulatory_applicability']['returned_dict_keys']
    assert 'residency_compliant' in tools['get_regulatory_applicability']['returned_dict_keys']
    assert tools['request_evidence']['calls_factual_snapshot']
    boundary={'source_sha256':digest(path),'methods':tools,
              'interpretation':'The direct-file certification must not be generalized to the whole tool interface. request_evidence returns the full gate factual_snapshot for an allowed evidence request, even if the requested ID is a non-snapshot document. A snapshot-removal ablation must control this return path. get_regulatory_applicability separately exposes residency_compliant. The two named getter methods do not themselves expose duplicate_capability or audit_trail in this frozen source.'}
    p=output/'source_coverage_audit.json';summary=read_json(p);summary['tool_boundary']=boundary
    evaluator_tree=ast.parse((SOURCE/'evaluator.py').read_text(encoding='utf-8'))
    general=next(n for n in evaluator_tree.body if isinstance(n,ast.FunctionDef) and n.name=='evaluate_general')
    external=collections.defaultdict(set)
    for node in ast.walk(general):
        if isinstance(node,ast.Subscript) and isinstance(node.value,ast.Name) and node.value.id=='r' and isinstance(node.slice,ast.Constant):
            external['upstream_results.[*].'+node.slice.value].add(node.lineno)
    summary['external_runtime_inputs']={'phase':'Supplied by the route occurrence, not a factual scalar in an evidence artifact.',
                                        'general_upstream':[{'field':k,'evaluator_lines':sorted(v)} for k,v in sorted(external.items())],
                                        'scope':'These runtime inputs are inventoried separately from the 76 case-fact/output-metadata fields; generated files cannot provide future agent handoff outcomes.'}
    p.write_text(json.dumps(summary,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    p=output/'source_coverage_audit.md';base=p.read_text(encoding='utf-8').split('\n## Tool-interface boundary\n')[0]
    p.write_text(base+'\n## Tool-interface boundary\n\n'+boundary['interpretation']+
                 '\n\nThe 76-field count concerns case facts and the output metadata `risk_owner`. Phase comes from the route manifest. General also consumes runtime `upstream_results[*].disposition` and `upstream_results[*].gate`; those AST dependencies are separately recorded in the JSON and are not expected in static source documents.\n',encoding='utf-8')


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--dataset',type=Path,default=ROOT/'experiments/preflight_balanced_300_20260922/dataset')
    ap.add_argument('--output-dir',type=Path,default=HERE)
    args=ap.parse_args(); dataset=args.dataset.resolve();output=args.output_dir.resolve()
    assert not output.is_relative_to(dataset)
    fields=inventory(); mapping=decoders()
    field_names={x['field'] for group in fields.values() for x in group}
    assert not set(mapping)-field_names, set(mapping)-field_names
    uncoded=sorted(field_names-set(mapping))
    cases=sorted(dataset.glob('DGF-*'));assert len(cases)==300
    requirements=[];occurrence_counts=collections.Counter();physical_counts=collections.Counter();per_field=collections.defaultdict(collections.Counter)
    source_hashes={};case_certified=[];occurrence_records=[]
    for folder in cases:
        reader=Sources(folder);route=read_json(folder/'01_route_manifest.json')['occurrences']
        case_fields={}
        for gate in GATES:
            snapshot=read_json(folder/f'gate_evidence/{gate}/review_facts.json')
            for item in fields[gate]:
                field=item['field'];expected=expected_value(snapshot,field); alternatives=[]
                for decoder in mapping.get(field,[]):
                    eid=decoder['evidence_id'];node=reader.nodes[eid];present=(folder/node['path']).is_file()
                    alt={'decoder':decoder,'physical_file_present':present,'path':node['path']}
                    try:
                        value,dependencies=reader.decode(decoder)
                        alt.update(decoded_value=value,snapshot_match=equal(value,expected),dependencies=dependencies,
                                   status='MATCH' if equal(value,expected) else 'MISMATCH')
                    except (KeyError,ValueError,OSError) as exc:
                        reason = 'MISSING_FILE' if isinstance(exc,FileNotFoundError) else str(exc)
                        alt.update(status='NOT_ESTABLISHED',reason=reason,dependencies=[eid])
                    alternatives.append(alt)
                matched=any(a['status']=='MATCH' for a in alternatives)
                physical='MATCH_ESTABLISHED' if matched else 'DECODED_VALUES_DISAGREE' if any(a['status']=='MISMATCH' for a in alternatives) else 'NOT_ESTABLISHED'
                case_fields[(gate,field)]={'case':folder.name,'gate':gate,'field':field,'role':item['role'],
                                          'snapshot_value':expected,'physical_status':physical,'alternatives':alternatives}
                physical_counts[physical]+=1;per_field[field][physical]+=1
                requirements.append(case_fields[(gate,field)])
        all_occ=True
        for occ in route:
            assessments=[]
            for item in fields[occ['gate']]:
                rec=case_fields[(occ['gate'],item['field'])];alts=[]
                for alt in rec['alternatives']:
                    access=[reader.availability(eid,occ['gate'],occ['phase']) for eid in alt['dependencies']]
                    alts.append({'evidence_id':alt['decoder']['evidence_id'],'status':alt['status'],
                                 'dependency_access':access,'readable_match':alt['status']=='MATCH' and all(x['direct_readable'] for x in access)})
                status='READABLE_MATCH_ESTABLISHED' if any(a['readable_match'] for a in alts) else 'PHYSICAL_MATCH_ACCESS_BLOCKED' if rec['physical_status']=='MATCH_ESTABLISHED' else rec['physical_status']
                occurrence_counts[status]+=1;per_field[item['field']]['occurrence_'+status]+=1
                assessments.append({'field':item['field'],'role':item['role'],'status':status,'alternatives':alts})
            complete=all(x['status']=='READABLE_MATCH_ESTABLISHED' for x in assessments);all_occ &= complete
            occurrence_records.append({'case':folder.name,'occurrence_id':occ['occurrence_id'],'gate':occ['gate'],'phase':occ['phase'],
                                       'all_static_fields_certified':complete,'fields':assessments})
        case_certified.append({'case':folder.name,'all_scheduled_occurrences_certified':all_occ})
        source_hashes[folder.name]=reader.hashes
        assert all(digest(folder/p)==h for p,h in reader.hashes.items()), 'Source changed during audit'
    proofs=omission_proofs(dataset)
    missing_files=sorted({(r['case'],a['path']) for r in requirements for a in r['alternatives'] if not a['physical_file_present']})
    nonmetadata_counts=collections.Counter(f['status'] for occ in occurrence_records for f in occ['fields'] if f['role']!='output_metadata')
    summary={
        'status':'CONSERVATIVE_CERTIFICATION_NOT_FULL_IDENTIFIABILITY_PROOF',
        'cases':len(cases),'scheduled_occurrences':len(occurrence_records),'unique_ast_fields':len(field_names),
        'gate_field_pairs':sum(len(x) for x in fields.values()),'mapped_unique_fields':len(mapping),'not_certified_fields':uncoded,
        'frozen_source_fingerprint':source_fingerprint(),'evaluator_sha256':digest(SOURCE/'evaluator.py'),
        'document_factory_sha256':digest(SOURCE/'document_factory.py'),
        'inventory':fields,'decoder_inventory':mapping,'physical_gate_field_counts':dict(physical_counts),
        'scheduled_occurrence_field_counts':dict(occurrence_counts),'per_field':{k:dict(v) for k,v in sorted(per_field.items())},
        'scheduled_nonmetadata_field_counts':dict(nonmetadata_counts),
        'unique_missing_candidate_source_files':len(missing_files),
        'fully_certified_scheduled_occurrences':sum(x['all_static_fields_certified'] for x in occurrence_records),
        'fully_certified_cases':sum(x['all_scheduled_occurrences_certified'] for x in case_certified),
        'generator_omission_counterexamples':proofs,
        'method':'AST inventory of all scalar accesses in each frozen evaluator, including the selected-vendor lookup. Labelled exact JSON/CSV and generated DOCX-table/clause decoders, never bare-value search. Public supplier selection comes from SCORING_MATRIX, not the snapshot. Decoded values are compared against the gate snapshot; metadata-only risk_owner is labelled separately. Direct read availability requires every decoder dependency to be in gate scope, in phase, AVAILABLE and physically present.',
        'limits':[
            'Static per-gate field inventory is an upper bound: it includes fields whose phase condition or short-circuited branch may not apply to a particular occurrence. Failure of all-static-fields certification does not prove that occurrence undecidable.',
            'A matching readable source is a coverage witness, not proof of unique inferability: conflicting or stale alternatives and authority resolution can still matter. Every alternative mismatch is retained.',
            'NOT_ESTABLISHED means no reliable implemented decoder or decoding failure; it is not a claim of physical absence. Exact scalar decoding intentionally does not substitute threshold predicates or undocumented semantic inference.',
            'Availability describes direct read_evidence at the frozen scope/phase/public_status. request_evidence, authoritative domain tools, extra prompt fields, diagram vision, and runtime UPSTREAM_DECISIONS are outside this certification and require their own controlled ablation interfaces.',
            'The selected-vendor offer decoder depends on SCORING_MATRIX and the source row jointly. A physically matching row is not declared readable when its selection dependency is blocked.',
            'DOCX conversion follows explicit frozen template labels and prose, not general document understanding. Source reader limits of 200 CSV rows/40000 DOCX characters are enforced.',
            'No independent human validation, model calls, original score changes, or new evidence generation in the dataset. Two counterfactual proofs capture generator payloads offline; they do not claim absence from every possible external tool.'
        ],'model_calls':0}
    details={'physical_requirements':requirements,'scheduled_occurrences':occurrence_records,'case_certification':case_certified,'read_source_sha256':source_hashes,
             'missing_candidate_source_files':[{'case':c,'path':p} for c,p in missing_files]}
    output.mkdir(parents=True,exist_ok=True)
    (output/'source_coverage_audit.json').write_text(json.dumps(summary,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    (output/'source_coverage_details.json').write_text(json.dumps(details,ensure_ascii=False,separators=(',',':'))+'\n',encoding='utf-8')
    lines=['# Non-snapshot source coverage: conservative offline certification','',
           f"All **{len(cases)} dossiers** and **{len(occurrence_records)} scheduled gates** were inspected. The AST inventory contains **{len(field_names)} distinct fields**, of which **{len(mapping)}** have explicit source decoders. All static gate/field pairs are inspected physically in every dossier, even where that gate is not scheduled.",'',
           'This is a certification inventory, not a model experiment or an assertion that every uncertified case is impossible. Static fields include phase-inactive branches. Matching sources may coexist with contradictions. Direct-read scope, phase, and availability are evaluated separately from physical existence.','',
           '| Physical gate/field status | Count |','|---|---:|']
    lines += [f'| {k} | {v} |' for k,v in sorted(physical_counts.items())]
    lines += ['', '| Scheduled gate/field status | Count |','|---|---:|']
    lines += [f'| {k} | {v} |' for k,v in sorted(occurrence_counts.items())]
    lines += ['', '## Uncertified fields','']+[f'- `{p}`: NOT_ESTABLISHED by the explicit decoders.' for p in uncoded]
    lines += ['', '## Proven generator omissions','']
    for p in proofs:
        lines += [f"- `{p['field']}`: on `{p['base_case']}`, changing `{p['variants'][0]['value']}` to `{p['variants'][1]['value']}` changes the frozen gate decision from GO to REWORK. Only `{p['changed_payloads'][0]}` changes; all {p['identical_non_snapshot_payloads']} non-snapshot output payloads and renderer inputs remain identical."]
    lines += ['', 'These counterexamples cover generated artifact contents and diagram inputs. They do not exclude a separate authoritative tool supplying the missing fact; such access would need to be specified explicitly in an ablation.','','## Field inventory','', '| Gate | AST fields |','|---|---|']
    lines += [f"| {g} | "+', '.join('`'+x['field']+'`' for x in fields[g])+' |' for g in GATES]
    lines += ['', '## Limits','']+['- '+x for x in summary['limits']]
    lines += ['', '## Reproduce','', '`python research/2026-09-followup/source_coverage_audit.py`', '',
              'The summary JSON contains AST line references, every decoder specification, counts and counterexamples. The detailed JSON preserves every comparison, decoded value, dependency access check and SHA-256 of read non-snapshot source files. No raw hidden truth is included in either output.']
    (output/'source_coverage_audit.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    annotate_tool_boundary(output)
    print(json.dumps({k:summary[k] for k in ['cases','scheduled_occurrences','unique_ast_fields','gate_field_pairs','mapped_unique_fields','not_certified_fields','physical_gate_field_counts','scheduled_occurrence_field_counts','fully_certified_scheduled_occurrences','fully_certified_cases','generator_omission_counterexamples']},indent=2))


if __name__=='__main__': main()
