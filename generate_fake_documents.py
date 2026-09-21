#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import random
import shutil
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Any

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from dgf_generator import generate_case, gate_decision
from azure_architecture import generate_architecture_profile, render_architecture, write_profile, architecture_flows
from semantic_profiles import apply_overrides, compute_security_rows, platform_readiness_rows

GATE_ORDER = [
    "general", "it", "architecture", "security", "tech_readiness",
    "procurement", "legal", "compliance"
]

GATE_LABELS = {
    "general": "Gate Général - Arbitrage Stratégique & Budgétaire",
    "it": "Gate IT - Gouvernance Technologique & Standards",
    "architecture": "Gate Architecture - Urbanisation & Contraintes Réseau",
    "security": "Gate Sécurité Architecture - Cybersécurité & Risques",
    "tech_readiness": "Gate Tech Readiness - Maturité & Fiabilité Technique",
    "procurement": "Gate Procurement - Achats, Contrats & Fournisseurs",
    "legal": "Gate Légale - Conditions Contractuelles & Responsabilités",
    "compliance": "Gate Compliance - Normes, Réglementations & Éthique",
}

GATE_OBJECTIVES = {
    "general": "Décider si le projet est stratégiquement justifié, finançable, réalisable et acceptable au regard des risques et bénéfices attendus.",
    "it": "Vérifier l'alignement technologique, l'exploitabilité, les capacités, les standards IT et l'impact sur le SI existant.",
    "architecture": "Valider l'architecture cible, l'urbanisation, les flux, les contraintes réseau, la résilience, la performance et la réversibilité.",
    "security": "Vérifier les contrôles de sécurité, l'exposition, les identités, les données, la journalisation, les vulnérabilités et les risques résiduels.",
    "tech_readiness": "Vérifier que la solution est suffisamment testée, observable, exploitable, restaurable et prête pour une mise en production contrôlée.",
    "procurement": "Évaluer la stratégie d'achat, les fournisseurs, le coût total, les SLA, les sous-traitants, la réversibilité et l'état de contractualisation.",
    "legal": "Vérifier les clauses contractuelles, la responsabilité, la propriété intellectuelle, la confidentialité, les droits d'audit et les obligations de sortie.",
    "compliance": "Vérifier les obligations réglementaires, la protection des données, la traçabilité, la conservation, les tiers et les exigences de conformité applicables.",
}

FAKE_FILES = {
    "general": ["Business_Case_v4.xlsx", "Portfolio_Prioritization_2027.xlsx", "Risk_Register_v6.xlsx", "Benefits_Realization_Plan.docx"],
    "it": ["IT_Standards_Assessment.xlsx", "CMDB_Extract.csv", "Run_Model_v3.docx", "Capacity_Assessment.xlsx", "Lifecycle_Roadmap.pptx"],
    "architecture": ["02_architecture_evidence/Architecture_Diagram_Detailed.png", "02_architecture_evidence/Architecture_Facts.json", "HLD_v5.pptx", "Network_Flow_Matrix.xlsx", "Data_Model_v2.pdf", "Performance_Test_Plan.docx", "Exit_Strategy.docx"],
    "security": ["Security_Assessment_v4.docx", "02_architecture_evidence/Architecture_Diagram_Detailed.png", "01_HLD_Architecture_Overview.docx", "02_Security_Architecture_Design.docx", "03_IAM_and_Trust_Boundaries.docx", "04_Security_Logging_and_Data_Flows.docx", "PenTest_Report.pdf", "Risk_Cards.xlsx"],
    "tech_readiness": ["01_Infrastructure_Resilience_Design.docx", "02_Backup_Management_Plan.docx", "03_Disaster_Recovery_Plan.docx", "04_Recovery_Test_Report.docx", "05_Observability_and_SIEM_Readiness.docx", "06_Production_Runbook_and_Handover.docx", "07_Capacity_and_Performance_Test_Report.docx", "GoLive_Checklist.xlsx"],
    "procurement": ["RFP_Scorecard.xlsx", "Commercial_Proposal.pdf", "Vendor_Due_Diligence.docx", "SLA_Draft.docx", "Exit_Assistance_Proposal.pdf"],
    "legal": ["MSA_Draft_v7.docx", "DPA_Draft_v3.docx", "Liability_Redlines.docx", "Audit_Clause_Review.docx", "Termination_Schedule.xlsx"],
    "compliance": ["DPIA_Draft.docx", "Regulatory_Mapping.xlsx", "Retention_Schedule.xlsx", "Compliance_Assessment.docx", "Audit_Evidence_Index.xlsx"],
}

COMMENT_AUTHORS = {
    "general": ["Sponsor", "PMO", "Finance", "Risk", "Business Owner"],
    "it": ["DSI", "IT Operations", "Platform Engineering", "FinOps", "ITSM"],
    "architecture": ["Chief Architect", "Solution Architect", "Network Architect", "Data Architect", "Platform Architect"],
    "security": ["Security Architect", "CISO", "IAM Lead", "SOC Lead", "Cyber Risk"],
    "tech_readiness": ["SRE Lead", "QA Lead", "Release Manager", "Production Manager", "Service Owner"],
    "procurement": ["Category Manager", "Buyer", "Vendor Manager", "Procurement Director", "Finance"],
    "legal": ["Legal Counsel", "Contract Manager", "Privacy Counsel", "IP Counsel", "Legal Director"],
    "compliance": ["Compliance Officer", "DPO", "GRC", "Internal Audit", "Regulatory Affairs"],
}

MISLEADING_COMMENTS = [
    "Le point devrait être acceptable en pratique, même si la preuve n'est pas encore jointe.",
    "Le fournisseur confirme oralement que ce contrôle est couvert ; le document suivra plus tard.",
    "Le planning ne permet pas d'attendre une nouvelle validation, à traiter après le go-live.",
    "Cela fonctionne déjà en environnement de test, je propose de considérer le point comme clos.",
    "Le risque paraît faible selon l'équipe projet ; pas besoin de bloquer pour ce sujet.",
]

STATUS_COMMENTS = {
    "PASS": [
        "Preuve reçue et cohérente avec le dossier.",
        "Point vérifié, aucune réserve majeure identifiée.",
        "Contrôle conforme aux critères annoncés.",
    ],
    "WARN": [
        "Acceptable sous réserve d'une clarification et d'une action datée.",
        "Le point est partiellement démontré ; une preuve complémentaire est demandée.",
        "La solution peut avancer, mais la réserve doit rester ouverte jusqu'à vérification.",
    ],
    "FAIL": [
        "Point bloquant : la condition n'est pas satisfaite dans l'état actuel du dossier.",
        "Écart majeur ; retour au projet requis avant avis favorable.",
        "Le risque / l'écart dépasse la tolérance prévue pour cette gate.",
    ],
    "MISSING": [
        "Preuve absente : impossible de conclure sans information complémentaire.",
        "Le document attendu n'a pas été fourni.",
        "L'équipe projet doit fournir la preuve avant décision définitive.",
    ],
    "NA": [
        "Non applicable au périmètre déclaré.",
        "Sujet hors périmètre pour ce cas.",
    ],
}

STATUS_COLORS = {
    "PASS": "C6EFCE",
    "WARN": "FFEB9C",
    "FAIL": "FFC7CE",
    "MISSING": "D9EAF7",
    "NA": "E7E6E6",
}

DECISION_COLORS = {"GREEN":"C6EFCE", "YELLOW":"FFEB9C", "RED":"FFC7CE"}

def _recalc_gate(gate):
    gate["decision"] = gate_decision(gate["subjects"])
    gate["summary"] = {
        "pass": sum(x["status"]=="PASS" for x in gate["subjects"]),
        "warn": sum(x["status"]=="WARN" for x in gate["subjects"]),
        "fail": sum(x["status"]=="FAIL" for x in gate["subjects"]),
        "missing": sum(x["status"]=="MISSING" for x in gate["subjects"]),
        "na": sum(x["status"]=="NA" for x in gate["subjects"]),
    }


def synchronize_case_with_architecture(case, profile, seed:int, conflict_rate:float=.15):
    """Make most gate facts coherent with the generated Azure HLD, while preserving
    a controlled share of deliberate contradictions for benchmark reasoning."""
    rng=random.Random(seed)
    def get(gate_key, topic):
        for x in case["gates"][gate_key]["subjects"]:
            if x["topic"]==topic: return x
        return None
    def maybe_set(sub, value, status=None):
        if sub is None: return
        # Intentional cross-document inconsistency: keep previous random content.
        if rng.random() < conflict_rate: return
        sub["value"] = value
        if status: sub["status"] = status

    # Architecture gate
    maybe_set(get("architecture","Urbanisation & API"), {
        "integration_style": "REST/Event-driven" if profile.get("api_management") or profile.get("messaging") else "REST",
        "api_gateway": bool(profile.get("api_management")),
        "domain_alignment": rng.choice(["Strong","Strong","Partial"])
    })
    maybe_set(get("architecture","Architecture applicative"), {
        "pattern": profile.get("compute_profile"), "coupling": rng.choice(["Low","Medium"]), "reference_pattern_compliant": True
    }, "PASS" if profile.get("compute_profile")!="vm_single" else "WARN")
    data_model={"transactional":"Relational","enterprise_sql":"Relational","postgres":"Relational","event_driven":"Event stream","analytics":"Lakehouse","ai":"Mixed"}.get(profile.get("data_profile"),"Mixed")
    maybe_set(get("architecture","Architecture de données"), {
        "model":data_model,"master_data_owner":case["project"]["sponsor"],"data_lineage":rng.choice(["Complete","Partial"])
    })
    conn="ExpressRoute" if profile.get("expressroute") else ("VPN" if profile.get("vpn") else ("Private Link" if profile.get("private_endpoints") else "Internet"))
    maybe_set(get("architecture","Connectivité Cloud & Inter-sites"), {
        "connectivity":conn,"sites":rng.randint(1,8),"segmented":profile.get("network_profile")!="single_vnet"
    }, "PASS" if profile.get("network_profile")!="single_vnet" else "WARN")
    maybe_set(get("architecture","Disponibilité & Répartition"), {
        "sla_pct":rng.choice([99.9,99.95,99.99]) if profile.get("multi_az") else rng.choice([99.0,99.5,99.9]),
        "zones":3 if profile.get("multi_az") else 1,
        "load_balanced":profile.get("compute_profile") not in ("vm_single","functions"),
        "failover_tested":profile.get("dr_tested") if profile.get("multi_region") else profile.get("restore_tested")
    }, "PASS" if profile.get("multi_az") and (not profile.get("multi_region") or profile.get("dr_tested")) else "WARN")

    # Security gate
    maybe_set(get("security","Surface d'exposition & segmentation"), {
        "internet_exposed":profile.get("internet_facing"),
        "segments": 4 if profile.get("network_profile") in ("hub_spoke","virtual_wan") else 1,
        "dmz_or_private_zone":profile.get("edge_pattern") not in ("frontdoor_only",) and (profile.get("private_endpoints") or not profile.get("internet_facing"))
    }, "PASS" if (not profile.get("internet_facing") or "WAF" in " ".join(profile.get("edge_services",[]))) else "FAIL")
    maybe_set(get("security","IAM & privilèges"), {
        "sso":True,"mfa":"Mandatory" if profile.get("conditional_access") else "Optional",
        "privileged_roles":rng.randint(2,12),
        "service_identity":"Managed identity" if profile.get("managed_identity") else "Service principal"
    }, "PASS" if profile.get("managed_identity") and profile.get("conditional_access") else "WARN")
    maybe_set(get("security","Logs, monitoring & détection"), {
        "logs_to_siem":profile.get("logs_to_siem"),"retention_days":rng.choice([90,180,365,730]),
        "alerting":"Operational" if profile.get("sentinel") else "Partial"
    }, "PASS" if profile.get("logs_to_siem") else "FAIL")
    maybe_set(get("security","Sauvegarde, reprise & résilience"), {
        "backup":profile.get("backup_enabled"),"restore_tested":profile.get("restore_tested"),
        "ransomware_isolation":bool(profile.get("backup_enabled") and rng.random()<.75)
    }, "PASS" if profile.get("backup_enabled") and profile.get("restore_tested") else "WARN")

    # Tech Readiness main gate
    maybe_set(get("tech_readiness","Observabilité"), {
        "metrics":True,"logs":True,"traces":bool(profile.get("application_insights")),
        "dashboard":"Ready" if profile.get("sentinel") or profile.get("application_insights") else "Partial"
    }, "PASS" if profile.get("logs_to_siem") else "WARN")
    maybe_set(get("tech_readiness","Sauvegarde & restauration"), {
        "backup_enabled":profile.get("backup_enabled"),"restore_tested":profile.get("restore_tested"),
        "last_restore_days":rng.choice([7,30,90]) if profile.get("restore_tested") else None
    }, "PASS" if profile.get("backup_enabled") and profile.get("restore_tested") else "FAIL")
    maybe_set(get("tech_readiness","Aptitude à la production"), {
        "open_blockers":sum(1 for d in profile.get("defects",[]) if d.get("severity") in ("high","critical")),
        "open_major_actions":len(profile.get("defects",[])),"go_live_owner":case["project"]["project_manager"]
    }, "PASS" if not any(d.get("severity") in ("high","critical") for d in profile.get("defects",[])) else "WARN")

    for k in ("architecture","security","tech_readiness"):
        _recalc_gate(case["gates"][k])
    case["overall_decision"] = "RED" if any(g["decision"]=="RED" for g in case["gates"].values()) else ("YELLOW" if any(g["decision"]=="YELLOW" for g in case["gates"].values()) else "GREEN")
    return case


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_text_color(cell, color):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(color)


def add_doc_header(doc: Document, label: str, case: Dict[str, Any], gate: Dict[str, Any]):
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CONFIDENTIEL - DONNÉES SYNTHÉTIQUES DGF-BENCH")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(180, 0, 0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(31, 78, 121)

    info = doc.add_table(rows=4, cols=4)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.style = "Table Grid"
    pairs = [
        ("Projet", case["project"]["project_name"]),
        ("Project ID", case["project"]["project_id"]),
        ("Type", case["project"]["project_type"]),
        ("Business Unit", case["project"]["business_unit"]),
        ("Sponsor", case["project"]["sponsor"]),
        ("Project Manager", case["project"]["project_manager"]),
        ("Criticité", case["project"]["business_criticality"]),
        ("Classification", case["project"]["data_classification"]),
    ]
    idx=0
    for row in info.rows:
        for pair_i in range(2):
            key,val = pairs[idx]
            c1 = row.cells[pair_i*2]
            c2 = row.cells[pair_i*2+1]
            c1.text = key
            c2.text = str(val)
            set_cell_shading(c1, "D9EAF7")
            for rr in c1.paragraphs[0].runs:
                rr.bold=True
            idx += 1

    p = doc.add_paragraph()
    r = p.add_run("Objectif de la gate : ")
    r.bold = True
    p.add_run(GATE_OBJECTIVES[gate["gate_key"]])

    p = doc.add_paragraph()
    r = p.add_run("Population / reviewer : ")
    r.bold = True
    p.add_run(f"{gate['population_owner']} / {gate['reviewer']}")


def format_value(v):
    if isinstance(v, dict):
        parts=[]
        for k,val in v.items():
            if isinstance(val, list):
                val = ", ".join(map(str,val))
            parts.append(f"{k}: {val}")
        return "\n".join(parts)
    if isinstance(v, list):
        return ", ".join(map(str,v))
    return str(v)


def generate_comment(rng, gate_key, subject):
    author = rng.choice(COMMENT_AUTHORS[gate_key])
    # 20% of WARN/FAIL/MISSING comments are intentionally optimistic/misleading.
    misleading = subject["status"] in ("WARN","FAIL","MISSING") and rng.random() < 0.20
    if misleading:
        body = rng.choice(MISLEADING_COMMENTS)
        truth = "misleading_or_unverified"
    else:
        body = rng.choice(STATUS_COMMENTS[subject["status"]])
        truth = "aligned_with_ground_truth"
    return {
        "author": author,
        "date": (datetime.now() - timedelta(days=rng.randint(0,30))).strftime("%Y-%m-%d"),
        "text": body,
        "truthfulness": truth,
    }




def _topic_map(case, gate_key):
    return {x["topic"]: x for x in case["gates"][gate_key]["subjects"]}


def _doc_setup(title, subtitle, case, color="1F4E79"):
    doc = Document()
    doc.styles['Normal'].font.name = 'Aptos'
    doc.styles['Normal'].font.size = Pt(9)
    doc.styles['Normal'].paragraph_format.space_after = Pt(3)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.6)
    sec.right_margin = Inches(0.6)
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("SYNTHETIC EVIDENCE - DGF-BENCH")
    r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(180,0,0)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(title); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor.from_string(color)
    if subtitle:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=p.add_run(subtitle); rr.italic=True; rr.font.size=Pt(10)
    t=doc.add_table(rows=2, cols=4); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    pairs=[("Project",case['project']['project_name']),("Project ID",case['project']['project_id']),
           ("Type",case['project']['project_type']),("Classification",case['project']['data_classification'])]
    for i,(k,v) in enumerate(pairs):
        row=i//2; col=(i%2)*2
        t.rows[row].cells[col].text=k; t.rows[row].cells[col+1].text=str(v)
        set_cell_shading(t.rows[row].cells[col], 'D9EAF7')
        for rr in t.rows[row].cells[col].paragraphs[0].runs: rr.bold=True
    return doc


def _add_section_table(doc, title, headers, rows, widths=None):
    doc.add_heading(title, level=1)
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=str(h); set_cell_shading(t.rows[0].cells[i], '1F4E79'); set_cell_text_color(t.rows[0].cells[i], 'FFFFFF')
        for rr in t.rows[0].cells[i].paragraphs[0].runs: rr.bold=True; rr.font.size=Pt(8)
    for vals in rows:
        c=t.add_row().cells
        for i,v in enumerate(vals):
            c[i].text=str(v)
            c[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            for pp in c[i].paragraphs:
                for rr in pp.runs: rr.font.size=Pt(8)
    return t


def _add_fake_comments(doc, rng, authors, topic, count=3):
    doc.add_heading('Reviewer / project comments', level=1)
    pool = [
        'Please confirm this against the authoritative system before closure.',
        'The current design is acceptable only if the stated evidence is actually attached.',
        'This point is expected to be resolved before production readiness.',
        'Project team believes the risk is low; reviewer requests objective evidence.',
        'A verbal confirmation is not considered sufficient evidence for this control.',
        'If the dependency changes, this assessment must be repeated.',
    ]
    for _ in range(count):
        p=doc.add_paragraph(style='List Bullet')
        p.add_run(f"{rng.choice(authors)} - {topic}: ").bold=True
        p.add_run(rng.choice(pool))


def _save_support_doc(doc, path):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('Synthetic document for benchmark use only - no real organization or system.')
    r.italic=True; r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)
    doc.save(path)
    return path


def build_security_support_docs(case, evidence_dir:Path, seed:int):
    rng=random.Random(seed+41000)
    evidence_dir.mkdir(parents=True, exist_ok=True)
    sec=_topic_map(case,'security'); arch=_topic_map(case,'architecture'); it=_topic_map(case,'it')
    outputs=[]; manifest={}

    # 1 - HLD
    doc=_doc_setup('High Level Design (HLD)', 'Logical architecture, trust boundaries, flows and non-functional assumptions', case)
    doc.add_heading('1. Scope and architecture context', level=1)
    p=doc.add_paragraph(); p.add_run('Purpose: ').bold=True; p.add_run('Describe the target logical architecture used as input to the Security & Architecture Gate.')
    profile=case.get('architecture_profile',{})
    diagram_png=profile.get('diagram_png')
    if diagram_png and Path(diagram_png).exists():
        doc.add_heading('2. Detailed Azure architecture diagram', level=1)
        doc.add_picture(str(diagram_png), width=Inches(7.15))
        cp=doc.add_paragraph('Synthetic Azure HLD generated from the same benchmark seed. The diagram is evidence to review, not ground truth by itself.')
        cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for rr in cp.runs: rr.italic=True; rr.font.size=Pt(8)
        _add_section_table(doc,'2.1 Azure architecture facts',['Field','Value'],[
            ('Edge path',' -> '.join(profile.get('edge_services',[]))),
            ('Compute profile',profile.get('compute_profile')),
            ('Data profile',profile.get('data_profile')),
            ('Network profile',profile.get('network_profile')),
            ('Primary region',profile.get('primary_region')),
            ('Secondary region',profile.get('secondary_region') or 'None'),
            ('Multi-AZ',profile.get('multi_az')),
            ('Multi-region',profile.get('multi_region')),
            ('RTO / RPO',f"{profile.get('rto_hours')} h / {profile.get('rpo_minutes')} min"),
            ('Private endpoints',profile.get('private_endpoints')),
            ('Logs to SIEM',profile.get('logs_to_siem')),
        ])
    components=[
        ('Users / Admins','Client','Internet / Corporate','OIDC / SSO','External entry point'),
        ('Web Front End','Application','DMZ / App Zone','TLS 443','User-facing presentation layer'),
        ('API / Service Layer','Application','App Zone','mTLS / TLS','Business logic and integrations'),
        ('Primary Data Store','Data','Data Zone','TLS / private endpoint','Confidential application data'),
        ('Identity Provider','Identity','Shared Services','OIDC / SAML','Authentication and conditional access'),
        ('SIEM / Monitoring','Security','Shared Services','HTTPS / syslog','Logs, alerts and detection'),
    ]
    if case['project']['architecture_pattern']=='Agentic AI':
        components.insert(3,('AI Orchestration Service','AI / Integration','App Zone','HTTPS','Agent orchestration and tool calls'))
    if profile:
        components=[]
        for svc in profile.get('edge_services',[]): components.append((svc,'Edge / WAF / CDN','Global Edge','HTTPS 443','Global ingress / filtering / content delivery'))
        for svc in profile.get('compute_services',[]): components.append((svc,'Compute','Application Spoke','Private / HTTPS','Application workload'))
        for svc in profile.get('data_services',[]): components.append((svc,'Data','Data Spoke','Private Link / TLS' if profile.get('private_endpoints') else 'TLS / service endpoint','Persistent data / messaging'))
        if profile.get('azure_firewall'): components.append(('Azure Firewall','Network Security','Hub VNet','L3-L7','Central inspection / egress control'))
        if profile.get('bastion'): components.append(('Azure Bastion','Management','Hub VNet','HTTPS','Administrative access'))
        components += [('Microsoft Entra ID','Identity','Shared Service','OIDC/SAML','Authentication / authorization'),('Log Analytics / Sentinel','Security','Shared Service','HTTPS','Logging / SIEM')]
    _add_section_table(doc,'3. Logical component inventory',['Component','Role','Zone','Protocol','Purpose'],components)
    # simple diagram
    doc.add_heading('3. Logical flow overview', level=1)
    diag=doc.add_table(rows=3, cols=5); diag.style='Table Grid'; diag.alignment=WD_TABLE_ALIGNMENT.CENTER
    labels=[['User','→','Web / API','→','Data Store'],['','↘','Identity','↙',''],['','→','SIEM / Monitoring','←','Admin / Ops']]
    for r_i,row in enumerate(labels):
        for c_i,val in enumerate(row):
            diag.rows[r_i].cells[c_i].text=val
            if val not in ('','→','←','↘','↙'):
                set_cell_shading(diag.rows[r_i].cells[c_i], rng.choice(['D9EAF7','E2F0D9','FCE4D6','E4DFEC']))
                for rr in diag.rows[r_i].cells[c_i].paragraphs[0].runs: rr.bold=True
    if profile:
        flows=[]
        for idx,f in enumerate(architecture_flows(profile),1):
            flows.append((f"F{idx:02d}",f['source'],f['destination'],f['protocol'],f['purpose'],case['project']['data_classification'] if f['trust']!='external' else 'External'))
    else:
        flows=[
            ('F01','User','Web Front End','HTTPS 443','User session / business data',case['project']['data_classification']),
            ('F02','Web Front End','API / Service Layer','HTTPS 443','API requests',case['project']['data_classification']),
            ('F03','API / Service Layer','Primary Data Store','TLS / DB protocol','Application records',case['project']['data_classification']),
            ('F04','API / Service Layer','Identity Provider','OIDC / HTTPS','Tokens / identity attributes','Confidential'),
            ('F05','Application components','SIEM / Monitoring','HTTPS / agent','Security and operations logs','Internal'),
        ]
    _add_section_table(doc,'4. Principal data flows',['ID','Source','Destination','Protocol','Data','Classification'],flows)
    ha=arch.get('Disponibilité & Répartition',{}).get('value',{})
    cont=it.get('Continuité de service',{}).get('value',{})
    _add_section_table(doc,'5. Key non-functional assumptions',['Area','Assumption'],[
        ('Availability',f"Azure profile multi-AZ={profile.get('multi_az','N/A')}; resilience={profile.get('resilience_profile','N/A')}; failover tested={profile.get('dr_tested',ha.get('failover_tested','N/A'))}"),
        ('Recovery',f"RTO={profile.get('rto_hours',cont.get('rto_hours','N/A'))} h; RPO={profile.get('rpo_minutes',cont.get('rpo_minutes','N/A'))} min; DR mode={profile.get('dr_mode','N/A')}"),
        ('Hosting',profile.get('compute_profile',case['project']['hosting_model'])),('Primary region',profile.get('primary_region',case['project']['primary_region'])),
        ('Data classification',case['project']['data_classification'])])
    _add_fake_comments(doc,rng,COMMENT_AUTHORS['security'],'HLD',4)
    path=evidence_dir/'01_HLD_Architecture_Overview.docx'; outputs.append(_save_support_doc(doc,path))
    manifest[path.name]={'type':'HLD','facts':{'flows':flows,'components':[x[0] for x in components]}}

    # 2 - Security Architecture Design
    doc=_doc_setup('Security Architecture Design', 'Security zones, IAM, cryptography, exposure and control decisions', case, color='C00000')
    profile=case.get('architecture_profile',{})
    exposure=sec['Surface d\'exposition & segmentation']['value']; iam=sec['IAM & privilèges']['value']; crypt=sec['Chiffrement']['value']; logs=sec['Logs, monitoring & détection']['value']
    if profile:
        _add_section_table(doc,'0. Azure security design summary',['Control','Configured value'],[
            ('Edge/WAF/CDN',' -> '.join(profile.get('edge_services',[]))),('Azure Firewall',profile.get('azure_firewall')),('DDoS Protection',profile.get('ddos_protection')),
            ('Private Endpoints',profile.get('private_endpoints')),('Managed Identity',profile.get('managed_identity')),('PIM',profile.get('pim')),('Conditional Access',profile.get('conditional_access')),
            ('Sentinel',profile.get('sentinel')),('Defender for Cloud',profile.get('defender_for_cloud')),('Logs to SIEM',profile.get('logs_to_siem'))])
    _add_section_table(doc,'1. Trust zones and exposure',['Control area','Design statement','Evidence status'],[
        ('Internet exposure',f"internet_exposed={exposure['internet_exposed']}",sec['Surface d\'exposition & segmentation']['status']),
        ('Segmentation',f"segments={exposure['segments']}; private/DMZ={exposure['dmz_or_private_zone']}",sec['Surface d\'exposition & segmentation']['status']),
        ('Administrative path','Administration through managed identity / privileged access path','REVIEW REQUIRED'),
        ('Data zone','Data store reachable only from application/service tier','DESIGN INTENT'),
    ])
    _add_section_table(doc,'2. Identity and access model',['Topic','Value','Security note'],[
        ('SSO',iam.get('sso'),'Federated authentication where applicable'),('MFA',iam.get('mfa'),'Mandatory for privileged/admin paths'),
        ('Privileged roles',iam.get('privileged_roles'),'Privileged roles must be individually attributable'),('Service identity',iam.get('service_identity'),'No shared human identity for service execution')])
    _add_section_table(doc,'3. Cryptography and secret handling',['Topic','Value','Expected control'],[
        ('Encryption in transit',crypt.get('in_transit'),'TLS for every external and inter-zone sensitive flow'),
        ('Encryption at rest',crypt.get('at_rest'),'Managed encryption on persistent stores'),
        ('Customer managed keys',crypt.get('customer_managed_keys'),'Required only where policy / risk classification demands it'),
        ('Secrets','Synthetic Key Vault / secret manager','No secrets in code or configuration repositories')])
    _add_section_table(doc,'4. Logging and detection architecture',['Topic','Value','Comment'],[
        ('Logs to SIEM',logs.get('logs_to_siem'),'Application, identity and security control logs'),('Retention days',logs.get('retention_days'),'Subject to compliance retention'),
        ('Alerting',logs.get('alerting'),'High-risk events must create actionable alerts'),('Time synchronization','Required','UTC/NTP across sources')])
    _add_section_table(doc,'5. Open security decisions',['Decision','Status','Owner'],[
        ('Privileged access path approved',rng.choice(['Open','Conditional','Approved']),'Security Architect'),
        ('External exposure exception',rng.choice(['None','Open risk card','Approved exception']),'Cyber Risk'),
        ('Penetration test closure',rng.choice(['Complete','Partial','Pending']),'Application Owner'),
        ('Residual risk acceptance',rng.choice(['Not required','Pending','Approved']),'Business Risk Owner'),
    ])
    _add_fake_comments(doc,rng,COMMENT_AUTHORS['security'],'Security architecture',5)
    path=evidence_dir/'02_Security_Architecture_Design.docx'; outputs.append(_save_support_doc(doc,path))
    manifest[path.name]={'type':'security_architecture','facts':{'exposure':exposure,'iam':iam,'cryptography':crypt,'logging':logs}}

    # 3 - IAM and Trust Boundaries
    doc=_doc_setup('IAM and Trust Boundaries', 'Identities, privileges, service principals and authorization boundaries', case, color='C00000')
    identities=[
        ('End user','Human','SSO','Standard user',rng.choice(['Compliant','Review'])),
        ('Application service','Workload',str(iam.get('service_identity')),'Scoped application role',rng.choice(['Compliant','Review','Gap'])),
        ('Operations admin','Human','Privileged access','Admin / break-glass',rng.choice(['Compliant','Conditional'])),
        ('CI/CD runner','Workload','Federated workload identity','Deployment only',rng.choice(['Compliant','Review'])),
        ('Monitoring collector','Workload','Managed service identity','Read telemetry / write logs',rng.choice(['Compliant','Review'])),
    ]
    _add_section_table(doc,'1. Identity inventory',['Principal','Type','Authentication','Privilege','Assessment'],identities)
    _add_section_table(doc,'2. Authorization rules',['Rule','Enforcement'],[
        ('Least privilege','Access is limited to declared application and operational functions.'),('Separation of duties','Deployment identity cannot approve its own policy exception.'),
        ('Privileged access','Time-bound privileged roles; no standing shared administrator account.'),('Service-to-service','Explicit audience/scope and network path required.'),
        ('Break-glass','Separate emergency identity, monitored and reviewed after use.')])
    _add_fake_comments(doc,rng,['IAM Lead','Security Architect','Platform Engineering'],'IAM',4)
    path=evidence_dir/'03_IAM_and_Trust_Boundaries.docx'; outputs.append(_save_support_doc(doc,path))
    manifest[path.name]={'type':'iam','facts':{'identities':identities,'source_iam':iam}}

    # 4 - Logging and Data flows
    doc=_doc_setup('Security Logging and Data Flow Matrix', 'Security-relevant flows, log sources and expected detection coverage', case, color='C00000')
    flow_rows=[]
    for fid,src,dst,proto,data,clas in flows:
        flow_rows.append((fid,src,dst,proto,data,clas,rng.choice(['Allowed','Allowed with condition','Review'])))
    _add_section_table(doc,'1. Security data flow matrix',['ID','Source','Destination','Protocol','Data','Class.','Disposition'],flow_rows)
    log_rows=[
        ('Identity provider','Sign-in, conditional access, privilege events',rng.choice(['SIEM','Local only']),'High'),
        ('API / service','Authentication, authorization, application errors',rng.choice(['SIEM','SIEM','Local only']),'High'),
        ('Data store','Admin actions, backup/restore, failed access',rng.choice(['SIEM','Local only']),'Medium'),
        ('CI/CD','Deployments, artifact provenance, policy failures',rng.choice(['SIEM','SIEM','None']),'High'),
        ('Network / WAF','Connections, blocks, anomaly events',rng.choice(['SIEM','Local only']),'High'),
    ]
    _add_section_table(doc,'2. Log-source coverage',['Source','Events','Destination','Priority'],log_rows)
    _add_fake_comments(doc,rng,['SOC Lead','Security Architect','SRE Lead'],'Logging & flows',4)
    path=evidence_dir/'04_Security_Logging_and_Data_Flows.docx'; outputs.append(_save_support_doc(doc,path))
    manifest[path.name]={'type':'logging_flows','facts':{'flows':flow_rows,'log_sources':log_rows}}

    return outputs, manifest


def build_tech_readiness_support_docs(case, evidence_dir:Path, seed:int):
    rng=random.Random(seed+51000)
    evidence_dir.mkdir(parents=True, exist_ok=True)
    profile=case.get('architecture_profile',{})
    tr=_topic_map(case,'tech_readiness'); arch=_topic_map(case,'architecture'); it=_topic_map(case,'it'); sec=_topic_map(case,'security')
    outputs=[]; manifest={}

    # 1 resilience
    doc=_doc_setup('Infrastructure Resilience Design', 'High availability, failure domains, SPOF analysis and failover strategy', case, color='7030A0')
    ha=arch['Disponibilité & Répartition']['value']; cap=it['Capacité & Hébergement']['value']
    if profile:
        _add_section_table(doc,'0. Architecture resilience facts',['Field','Value'],[
            ('Compute',profile.get('compute_profile')),('Resilience profile',profile.get('resilience_profile')),('Primary region',profile.get('primary_region')),('Secondary region',profile.get('secondary_region') or 'None'),
            ('Availability zones',', '.join(map(str,profile.get('availability_zones',[])))),('Multi-region',profile.get('multi_region')),('Azure Site Recovery',profile.get('azure_site_recovery')),
            ('Data replication',profile.get('data_replication')),('RTO',f"{profile.get('rto_hours')} h"),('RPO',f"{profile.get('rpo_minutes')} min"),('DR tested',profile.get('dr_tested'))])
    resilience_rows=[
        ('Compute tier','Active/active' if ha.get('zones',1)>1 else 'Single zone',f"Zones: {ha.get('zones')}",rng.choice(['PASS','WARN'])),
        ('Load balancing','Application / platform load balancer',f"Configured={ha.get('load_balanced')}",rng.choice(['PASS','WARN'])),
        ('Data tier',rng.choice(['Zone redundant','Primary/replica','Single instance']),rng.choice(['Automatic failover','Manual failover','No tested failover']),rng.choice(['PASS','WARN','FAIL'])),
        ('Network path',rng.choice(['Dual path','Single VPN','ExpressRoute + VPN backup','Internet only']),rng.choice(['Tested quarterly','Not recently tested']),rng.choice(['PASS','WARN'])),
        ('Capacity headroom',f"{cap.get('headroom_pct')}%",'Target >= 20%', 'WARN' if int(cap.get('headroom_pct',0))<20 else 'PASS'),
        ('Single points of failure',rng.choice(['None identified','DNS dependency','Single database writer','Vendor API']),rng.choice(['Mitigated','Open action','Accepted']),rng.choice(['PASS','WARN','FAIL'])),
    ]
    _add_section_table(doc,'1. Resilience control matrix',['Layer','Design','Evidence','Assessment'],resilience_rows)
    _add_section_table(doc,'2. Failure scenarios and response',['Scenario','Expected behavior','Recovery action'],[
        ('Single node loss','Traffic remains available','Automatic replacement / scale-out'),('Availability zone loss','Service continues or fails over','Route traffic to surviving zone'),
        ('Database failure','Writes recover on replica','Promote replica / validate integrity'),('Identity dependency outage','Existing sessions handled per design','Fallback / graceful degradation'),
        ('Monitoring outage','Service remains operational but alerting degraded','Restore telemetry path and manual observation')])
    _add_fake_comments(doc,rng,['SRE Lead','Platform Engineering','Production Manager'],'Infrastructure resilience',4)
    path=evidence_dir/'01_Infrastructure_Resilience_Design.docx'; outputs.append(_save_support_doc(doc,path)); manifest[path.name]={'type':'resilience','facts':{'ha':ha,'capacity':cap,'rows':resilience_rows}}

    # 2 backup management
    doc=_doc_setup('Backup Management Plan', 'Backup scope, schedules, retention, immutability and restoration controls', case, color='7030A0')
    backup=tr['Sauvegarde & restauration']['value']
    backup_rows=[
        ('Primary database',rng.choice(['Every 4h','Hourly','Daily']),rng.choice(['35 days','90 days','1 year']),rng.choice(['Encrypted','Encrypted + immutable']),rng.choice(['Secondary region','Same region vault'])),
        ('Configuration / IaC','On every change',rng.choice(['180 days','1 year','Git history']),rng.choice(['Repository encryption','Signed commits']),rng.choice(['Git / artifact repository'])),
        ('Object / file data',rng.choice(['Daily','Every 12h']),rng.choice(['30 days','90 days','7 years']),rng.choice(['Encrypted','Immutable 14 days']),rng.choice(['Secondary region','Offline copy'])),
        ('Secrets / keys','Metadata backup / documented recovery',rng.choice(['Current + previous','Versioned']),rng.choice(['HSM / Key Vault protected']),rng.choice(['Managed service'])),
    ]
    _add_section_table(doc,'1. Backup policy',['Scope','Frequency','Retention','Protection','Copy location'],backup_rows)
    _add_section_table(doc,'2. Backup control status',['Control','Value'],[
        ('Backup enabled',backup.get('backup_enabled')),('Restore tested',backup.get('restore_tested')),('Last restore test days',backup.get('last_restore_days')),
        ('Monitoring of failed backup jobs',rng.choice(['Enabled','Partial','Missing'])),('Backup access model',rng.choice(['Dedicated backup operators','Platform admins','Vendor managed'])),
        ('Immutability',rng.choice(['Enabled','Partial','Not enabled']))])
    _add_fake_comments(doc,rng,['Backup Service Owner','SRE Lead','Security Architect'],'Backup management',4)
    path=evidence_dir/'02_Backup_Management_Plan.docx'; outputs.append(_save_support_doc(doc,path)); manifest[path.name]={'type':'backup','facts':{'backup':backup,'policies':backup_rows}}

    # 3 DR / recovery plan
    doc=_doc_setup('Disaster Recovery and Recovery Plan', 'RTO/RPO, invocation criteria, recovery sequence and responsibilities', case, color='7030A0')
    cont=it['Continuité de service']['value']
    rto=cont.get('rto_hours'); rpo=cont.get('rpo_minutes')
    _add_section_table(doc,'1. Recovery objectives',['Service','RTO','RPO','Business criticality'],[(case['project']['project_name'],f"{rto} h",f"{rpo} min",case['project']['business_criticality'])])
    _add_section_table(doc,'2. Invocation criteria',['Trigger','Decision / action'],[
        ('Primary region unavailable > 15 minutes','Invoke DR incident and assess failover'),('Database unrecoverable in primary','Promote validated replica or restore backup'),
        ('Cyber incident affecting integrity','Isolate, preserve evidence, recover from trusted restore point'),('Loss of critical third-party dependency','Activate degraded mode / manual fallback where defined')])
    recovery_steps=[
        ('1','Declare incident and freeze unsafe changes','Incident Commander'),('2','Validate latest clean recovery point','DB / Backup Lead'),('3','Restore core identity / secrets / platform dependencies','Platform Lead'),
        ('4','Restore data services','Database Lead'),('5','Restore application and integrations','Application Owner'),('6','Run technical and business validation','QA + Business Owner'),('7','Authorize return to service','Service Owner / Incident Commander')]
    _add_section_table(doc,'3. Recovery sequence',['Step','Activity','Owner'],recovery_steps)
    _add_section_table(doc,'4. Communications and dependencies',['Area','Plan'],[
        ('Stakeholders','Business owner, operations, security, service desk, vendor'),('Status cadence','Every 30 minutes during Sev1'),('Vendor escalation',rng.choice(['Tested','Documented','Incomplete'])),
        ('DNS / certificate dependencies',rng.choice(['Documented','Partial','Missing'])),('Return-to-primary plan',rng.choice(['Documented','Draft','Missing']))])
    _add_fake_comments(doc,rng,['BCM Lead','SRE Lead','Service Owner','Security Incident Manager'],'Recovery plan',4)
    path=evidence_dir/'03_Disaster_Recovery_Plan.docx'; outputs.append(_save_support_doc(doc,path)); manifest[path.name]={'type':'dr_plan','facts':{'rto_hours':rto,'rpo_minutes':rpo,'steps':recovery_steps}}

    # 4 recovery test report
    doc=_doc_setup('Recovery Test Report', 'Synthetic DR / restore exercise with measured results and gaps', case, color='7030A0')
    actual_rto=round(float(rto)*rng.uniform(0.6,1.8),2) if rto else round(rng.uniform(2,16),2)
    actual_rpo=round(float(rpo)*rng.uniform(0.5,2.0),1) if rpo else rng.choice([15,60,240])
    rto_pass=actual_rto<=float(rto) if rto else False; rpo_pass=actual_rpo<=float(rpo) if rpo else False
    _add_section_table(doc,'1. Test result',['Metric','Target','Measured','Result'],[
        ('RTO',f"{rto} h",f"{actual_rto} h",'PASS' if rto_pass else 'FAIL'),('RPO',f"{rpo} min",f"{actual_rpo} min",'PASS' if rpo_pass else 'FAIL'),
        ('Restore integrity','No corruption',rng.choice(['No corruption detected','Minor reconciliation issue','Validation incomplete']),rng.choice(['PASS','WARN'])),
        ('Application smoke test','All critical journeys',rng.choice(['100%','95%','80%']),rng.choice(['PASS','WARN','FAIL']))])
    gaps=[('Dependency startup ordering',rng.choice(['Closed','Open']),rng.choice(['Medium','High'])),('Manual DNS step',rng.choice(['Closed','Open']),rng.choice(['Medium','High'])),('Vendor contact availability',rng.choice(['Closed','Open']),rng.choice(['Low','Medium']))]
    _add_section_table(doc,'2. Observed gaps',['Gap','Status','Severity'],gaps)
    _add_fake_comments(doc,rng,['QA Lead','SRE Lead','BCM Lead'],'Recovery test',4)
    path=evidence_dir/'04_Recovery_Test_Report.docx'; outputs.append(_save_support_doc(doc,path)); manifest[path.name]={'type':'recovery_test','facts':{'actual_rto':actual_rto,'actual_rpo':actual_rpo,'rto_pass':rto_pass,'rpo_pass':rpo_pass,'gaps':gaps}}

    # 5 observability
    doc=_doc_setup('Observability and SIEM Readiness', 'Metrics, logs, traces, alerting, dashboards and on-call readiness', case, color='7030A0')
    obs=tr['Observabilité']['value']; logs=sec['Logs, monitoring & détection']['value']
    _add_section_table(doc,'1. Telemetry coverage',['Signal','Coverage','Destination','Assessment'],[
        ('Metrics',obs.get('metrics'),'Monitoring platform',rng.choice(['PASS','WARN'])),('Application logs',obs.get('logs'), 'SIEM' if logs.get('logs_to_siem') else 'Local only', 'PASS' if logs.get('logs_to_siem') else 'FAIL'),
        ('Distributed traces',obs.get('traces'),'APM / tracing',rng.choice(['PASS','WARN'])),('Security logs',logs.get('logs_to_siem'),'SIEM', 'PASS' if logs.get('logs_to_siem') else 'FAIL'),('Dashboard',obs.get('dashboard'),'Operations dashboard',rng.choice(['PASS','WARN']))])
    _add_section_table(doc,'2. Operational alerts',['Alert','Threshold','Routing'],[
        ('Availability','Synthetic check failed 3 times','On-call Sev2'),('Error rate','>5% for 5 minutes','Application on-call'),('Latency','p95 exceeds SLO','SRE'),('Authentication failures','Anomalous rate','SOC'),('Backup failure','Any critical backup failed','Backup owner + SRE')])
    _add_fake_comments(doc,rng,['SRE Lead','SOC Lead','Service Owner'],'Observability',4)
    path=evidence_dir/'05_Observability_and_SIEM_Readiness.docx'; outputs.append(_save_support_doc(doc,path)); manifest[path.name]={'type':'observability','facts':{'observability':obs,'logging':logs}}

    # 6 runbook / handover
    doc=_doc_setup('Production Runbook and Operations Handover', 'Operational procedures, escalation, on-call and ownership', case, color='7030A0')
    run=tr['Runbooks & exploitation']['value']; sup=tr['Support & incidents']['value']; dep=tr['Déploiement & rollback']['value']
    _add_section_table(doc,'1. Ownership and support',['Item','Value'],[
        ('Runbook status',run.get('runbooks')),('On-call defined',run.get('on_call')),('Handover signed',run.get('handover_signed')),('Support model',sup.get('support_model')),('Sev1 response',f"{sup.get('sev1_response_min')} min")])
    _add_section_table(doc,'2. Standard operating procedures',['Procedure','Summary'],[
        ('Start / stop','Controlled service start/stop with dependency checks'),('Deployment',f"Automated={dep.get('deployment_automated')}; change type={dep.get('change_type')}"),('Rollback',f"Tested={dep.get('rollback_tested')}"),
        ('Incident triage','Check monitoring, recent changes, dependencies, security indicators'),('Backup verification','Review last successful backup and restore-test status'),('Escalation','SRE -> Service Owner -> Security / Vendor according to symptom')])
    _add_fake_comments(doc,rng,['Production Manager','SRE Lead','Service Owner'],'Runbook / handover',4)
    path=evidence_dir/'06_Production_Runbook_and_Handover.docx'; outputs.append(_save_support_doc(doc,path)); manifest[path.name]={'type':'runbook_handover','facts':{'runbook':run,'support':sup,'deployment':dep}}

    # 7 capacity/performance
    doc=_doc_setup('Capacity and Performance Test Report', 'Load, latency, saturation and scale test evidence', case, color='7030A0')
    perf=tr['Performance & charge']['value']; aperf=arch['Performance & Scalabilité']['value']
    _add_section_table(doc,'1. Performance results',['Metric','Target / plan','Measured / status'],[
        ('Peak test',f"{aperf.get('peak_rps')} RPS planned",f"Tested at {perf.get('peak_test_pct')}% of peak"),('p95 latency','Application SLO',f"{perf.get('p95_ms')} ms"),
        ('Autoscaling','Expected for variable demand',aperf.get('autoscaling')),('Load test evidence','Required before go-live',perf.get('load_tested')),('Known capacity headroom','>=20% recommended',f"{it['Capacité & Hébergement']['value'].get('headroom_pct')}%")])
    _add_fake_comments(doc,rng,['QA Lead','Performance Engineer','SRE Lead'],'Capacity / performance',4)
    path=evidence_dir/'07_Capacity_and_Performance_Test_Report.docx'; outputs.append(_save_support_doc(doc,path)); manifest[path.name]={'type':'performance','facts':{'tech_readiness':perf,'architecture':aperf}}

    return outputs, manifest


def build_semantic_platform_docs(case, seed:int):
    """Create compute-specific security/readiness documents driven by the selected Azure platform."""
    profile=case.get("architecture_profile",{})
    rng=random.Random(seed+88000)
    created=[]; manifests={}
    sec_dir=Path(profile.get("case_security_dir","")) if profile.get("case_security_dir") else None
    tech_dir=Path(profile.get("case_tech_dir","")) if profile.get("case_tech_dir") else None
    if sec_dir:
        sec_dir.mkdir(parents=True,exist_ok=True)
        doc=_doc_setup("Compute Platform Security Design",
                       f"Security controls specific to {profile.get('compute_profile')} workloads",case,color="C00000")
        c,rows=compute_security_rows(profile)
        _add_section_table(doc,"1. Platform-specific security controls",["Control","Generated value"],rows)
        er=profile.get("semantic_edge",{})
        _add_section_table(doc,"2. Edge / WAF / CDN controls",["Control","Generated value"],[
            ("Pattern",er.get("pattern")),("Services"," -> ".join(er.get("services",[]))),
            ("WAF present",er.get("waf_present")),("WAF mode",er.get("waf_mode")),
            ("Managed rules",er.get("managed_rules")),("Bot protection",er.get("bot_protection")),
            ("Rate limiting",er.get("rate_limiting")),("Origin restriction",er.get("origin_restriction")),
            ("CDN TTL",er.get("cdn_cache_ttl_seconds")),("TLS minimum",er.get("tls_min")),
        ])
        nr=profile.get("semantic_network",{})
        _add_section_table(doc,"3. Network security controls",["Control","Generated value"],[(k.replace("_"," ").title(),str(v)) for k,v in nr.items()])
        faults=[f for f in profile.get("semantic_faults",[]) if f.get("layer") in ("compute","edge","data")]
        doc.add_heading("4. Reviewer findings",level=1)
        if faults:
            for f in faults:
                p=doc.add_paragraph(style="List Bullet"); p.add_run(f"[{f['severity'].upper()}] {f['id']}: ").bold=True; p.add_run(f["description"])
        else:
            doc.add_paragraph("No platform-specific hidden defect generated in this synthetic case.")
        _add_fake_comments(doc,rng,COMMENT_AUTHORS["security"],c,4)
        p=sec_dir/"05_Compute_Platform_Security_Design.docx"; created.append(_save_support_doc(doc,p))
        manifests[p.name]={"type":"compute_security","compute_profile":c,"rows":rows,"faults":faults}

    if tech_dir:
        tech_dir.mkdir(parents=True,exist_ok=True)
        doc=_doc_setup("Platform Operations & Readiness",
                       f"Operational readiness for {profile.get('compute_profile')} compute",case,color="7030A0")
        rows=platform_readiness_rows(profile)
        _add_section_table(doc,"1. Platform-specific production readiness",["Check","Generated value"],rows)
        dd=profile.get("semantic_data",{})
        _add_section_table(doc,"2. Data platform readiness",["Check","Generated value"],[(k.replace("_"," ").title(),str(v)) for k,v in dd.items()])
        _add_section_table(doc,"3. Resilience and recovery posture",["Check","Value"],[
            ("Resilience profile",profile.get("resilience_profile")),("Multi-AZ",profile.get("multi_az")),
            ("Multi-region",profile.get("multi_region")),("Primary region",profile.get("primary_region")),
            ("Secondary region",profile.get("secondary_region") or "None"),("RTO",f"{profile.get('rto_hours')} h"),
            ("RPO",f"{profile.get('rpo_minutes')} min"),("Backup enabled",profile.get("backup_enabled")),
            ("Restore tested",profile.get("restore_tested")),("DR tested",profile.get("dr_tested")),
            ("Data replication",profile.get("data_replication")),("ASR",profile.get("azure_site_recovery")),
        ])
        _add_fake_comments(doc,rng,COMMENT_AUTHORS["tech_readiness"],profile.get("compute_profile"),4)
        p=tech_dir/"08_Platform_Operations_and_Readiness.docx"; created.append(_save_support_doc(doc,p))
        manifests[p.name]={"type":"platform_readiness","compute_profile":profile.get("compute_profile"),"rows":rows}
    return created,manifests


def create_gate_doc(case, gate_key, out_dir, seed):
    gate = case["gates"][gate_key]
    rng = random.Random(seed + hash(gate_key)%10000)
    doc = Document()
    doc.styles['Normal'].font.name = 'Aptos'
    doc.styles['Normal'].font.size = Pt(8.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)
    doc.styles['Normal'].paragraph_format.space_before = Pt(0)
    add_doc_header(doc, GATE_LABELS[gate_key], case, gate)

    doc.add_heading("1. Documents et preuves annoncés", level=1)
    for f in FAKE_FILES[gate_key]:
        p=doc.add_paragraph(style=None)
        p.style = doc.styles['List Bullet']
        p.add_run(f)

    doc.add_heading("2. Points de contrôle et informations du dossier", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Sujet", "Information générée", "Statut", "Preuve / Référence", "Commentaire projet / reviewer", "Origine"]
    for i,h in enumerate(headers):
        table.rows[0].cells[i].text=h
        set_cell_shading(table.rows[0].cells[i], "1F4E79")
        set_cell_text_color(table.rows[0].cells[i], "FFFFFF")
        for r in table.rows[0].cells[i].paragraphs[0].runs:
            r.bold=True
            r.font.size=Pt(8)

    comment_log=[]
    for s in gate["subjects"]:
        c = generate_comment(rng, gate_key, s)
        comment_log.append({"topic":s["topic"], **c})
        row=table.add_row().cells
        row[0].text=s["topic"]
        row[1].text=format_value(s["value"])
        row[2].text=s["status"]
        row[3].text="\n".join(s["evidence"]) if s["evidence"] else "Aucune preuve jointe"
        row[4].text=f"[{c['date']}] {c['author']}: {c['text']}"
        row[5].text=s["topic_origin"]
        set_cell_shading(row[2], STATUS_COLORS[s["status"]])
        for cell in row:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size=Pt(7.5)

    doc.add_heading("3. Fil de commentaires", level=1)
    for item in comment_log[:min(6,len(comment_log))]:
        p=doc.add_paragraph()
        p.add_run(f"{item['date']} - {item['author']} - {item['topic']}: ").bold=True
        p.add_run(item['text'])

    # Cross-gate / project comments
    doc.add_heading("4. Commentaires transverses du projet", level=1)
    generic = [
        ("Project Manager", "La date cible est considérée comme prioritaire ; merci de signaler uniquement les vrais blockers."),
        ("Sponsor", "Le budget ne doit pas augmenter sans validation du comité."),
        ("PMO", "Toutes les actions jaunes doivent avoir un owner et une date avant la prochaine gate."),
    ]
    if rng.random()<0.5:
        generic.append(("Vendor", "Le fournisseur indique que les preuves complémentaires peuvent être partagées sous NDA."))
    for author, txt in generic:
        p=doc.add_paragraph(style='List Bullet')
        p.add_run(f"{author}: ").bold=True
        p.add_run(txt)

    doc.add_heading("5. Décision de gate", level=1)
    dec_table = doc.add_table(rows=2, cols=5)
    dec_table.style = "Table Grid"
    labels = ["Décision", "PASS", "WARN", "FAIL", "MISSING"]
    vals = [gate["decision"], gate["summary"]["pass"], gate["summary"]["warn"], gate["summary"]["fail"], gate["summary"]["missing"]]
    for i,l in enumerate(labels):
        dec_table.rows[0].cells[i].text=l
        set_cell_shading(dec_table.rows[0].cells[i], "D9EAF7")
        dec_table.rows[1].cells[i].text=str(vals[i])
    set_cell_shading(dec_table.rows[1].cells[0], DECISION_COLORS[gate["decision"]])

    open_actions=[]
    for s in gate["subjects"]:
        if s["status"] in ("WARN","FAIL","MISSING"):
            open_actions.append({
                "topic":s["topic"],
                "action": rng.choice([
                    "Fournir la preuve manquante et faire revalider le point.",
                    "Corriger l'écart puis joindre la nouvelle preuve.",
                    "Documenter l'exception, l'owner, la date et la décision d'acceptation.",
                    "Clarifier l'information contradictoire auprès de la source autoritative.",
                ]),
                "owner": rng.choice(COMMENT_AUTHORS[gate_key]),
                "due_date": (datetime.now()+timedelta(days=rng.randint(5,45))).strftime("%Y-%m-%d")
            })
    doc.add_heading("6. Actions ouvertes", level=1)
    if open_actions:
        omitted_actions = max(0, len(open_actions)-4)
        open_actions = open_actions[:4]
        act = doc.add_table(rows=1, cols=4)
        act.style="Table Grid"
        for i,h in enumerate(["Sujet","Action","Owner","Due date"]):
            act.rows[0].cells[i].text=h
            set_cell_shading(act.rows[0].cells[i], "1F4E79")
            set_cell_text_color(act.rows[0].cells[i], "FFFFFF")
        for a in open_actions:
            row=act.add_row().cells
            row[0].text=a["topic"]
            row[1].text=a["action"]
            row[2].text=a["owner"]
            row[3].text=a["due_date"]
    else:
        doc.add_paragraph("Aucune action ouverte.")

    doc.add_paragraph()
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Document entièrement synthétique - à utiliser uniquement pour tests et benchmark.")
    r.italic=True
    r.font.size=Pt(8)
    r.font.color.rgb=RGBColor(100,100,100)

    out_path = out_dir / f"{GATE_ORDER.index(gate_key)+1:02d}_{gate_key}_{gate['decision']}.docx"
    doc.save(out_path)
    return out_path, comment_log, open_actions


def build_case_pack(seed:int, difficulty:int, project_type:str|None, out_root:Path, conflict_rate:float, architecture_inconsistency_rate:float=0.15, architecture_overrides:dict|None=None):
    case = generate_case(seed=seed, difficulty=difficulty, project_type=project_type, gates=None, conflict_rate=conflict_rate)
    case_dir=out_root / f"{case['project']['project_id']}_{case['project']['project_type']}"
    case_dir.mkdir(parents=True,exist_ok=True)
    # Detailed random Azure HLD generated from the same case seed.
    arch_profile=generate_architecture_profile(seed+900000, case['project'], difficulty, architecture_inconsistency_rate)
    arch_profile=apply_overrides(arch_profile, architecture_overrides, seed+900100, difficulty)
    case=synchronize_case_with_architecture(case,arch_profile,seed+901000,architecture_inconsistency_rate)
    arch_dir=case_dir/'02_architecture_evidence'; arch_dir.mkdir(parents=True,exist_ok=True)
    svg=arch_dir/'Architecture_Diagram_Detailed.svg'; png=arch_dir/'Architecture_Diagram_Detailed.png'; facts=arch_dir/'Architecture_Facts.json'
    render_architecture(arch_profile,case['project'],svg,png); write_profile(arch_profile,case['project'],facts)
    arch_profile['diagram_svg']=str(svg); arch_profile['diagram_png']=str(png); arch_profile['facts_file']=str(facts)
    case['architecture_profile']=arch_profile
    case['architecture_profile']['case_security_dir']=str(case_dir/'04_security_evidence')
    case['architecture_profile']['case_tech_dir']=str(case_dir/'05_tech_readiness_evidence')
    docs=[]
    hidden={
        "case_id":case["case_id"],
        "project":case["project"],
        "overall_decision":case["overall_decision"],
        "gates":{},
        "cross_gate_conflicts":case["cross_gate_conflicts"],
        "architecture_profile":case["architecture_profile"],
        "architecture_flows":architecture_flows(case["architecture_profile"]),
        "benchmark_note":"NE PAS FOURNIR CE FICHIER A L'AGENT EVALUE. Il contient la vérité terrain, les défauts d'architecture et la nature de certains commentaires trompeurs."
    }
    for i,key in enumerate(GATE_ORDER):
        p, comments, actions = create_gate_doc(case,key,case_dir,seed+i*1000)
        docs.append(p)
        support_docs=[]
        support_manifest={}
        if key == "security":
            support_dir=case_dir / "04_security_evidence"
            support_docs, support_manifest = build_security_support_docs(case, support_dir, seed+i*1000)
            docs.extend(support_docs)
        elif key == "tech_readiness":
            support_dir=case_dir / "05_tech_readiness_evidence"
            support_docs, support_manifest = build_tech_readiness_support_docs(case, support_dir, seed+i*1000)
            docs.extend(support_docs)
        hidden["gates"][key]={
            "expected_decision":case["gates"][key]["decision"],
            "subjects":case["gates"][key]["subjects"],
            "comments":comments,
            "open_actions":actions,
            "support_documents":support_manifest,
        }
    semantic_docs, semantic_manifest = build_semantic_platform_docs(case, seed)
    docs.extend(semantic_docs)
    hidden["semantic_support_documents"] = semantic_manifest
    public_context=dict(case["project"]); public_context["architecture"]={k:v for k,v in case["architecture_profile"].items() if k not in ("defects","inconsistencies","semantic_faults","diagram_png","diagram_svg","facts_file","case_security_dir","case_tech_dir")}
    (case_dir/"00_project_context.json").write_text(json.dumps(public_context,ensure_ascii=False,indent=2),encoding="utf-8")
    (case_dir/"99_hidden_ground_truth.json").write_text(json.dumps(hidden,ensure_ascii=False,indent=2),encoding="utf-8")
    readme = [
        f"DGF synthetic case {case['project']['project_id']}",
        "",
        f"Project: {case['project']['project_name']}",
        f"Type: {case['project']['project_type']}",
        f"Difficulty: {difficulty}",
        f"Overall expected decision: {case['overall_decision']}",
        "",
        "Each DOCX is a realistic synthetic gate dossier with visible stakeholder comments.",
        "99_hidden_ground_truth.json must NOT be shown to the evaluated agent.",
    ]
    (case_dir/"README.txt").write_text("\n".join(readme),encoding="utf-8")
    return case_dir, docs


def main():
    ap=argparse.ArgumentParser(description="Generate realistic fake DGF gate documents (DOCX) for benchmarking.")
    ap.add_argument("--count",type=int,default=1)
    ap.add_argument("--seed",type=int,default=42)
    ap.add_argument("--difficulty",type=int,choices=range(1,6),default=3)
    ap.add_argument("--project-type",choices=["new_platform","ma_integration","new_project"],default=None)
    ap.add_argument("--conflict-rate",type=float,default=0.08,help="Cross-gate textual conflict rate")
    ap.add_argument("--architecture-inconsistency-rate",type=float,default=0.15,help="Chance that a hidden architecture defect is represented inconsistently across documents")
    ap.add_argument("--force-compute",choices=["vm_single","vm_ha","vmss","dedicated_host","aks","container_apps","app_service","functions","mixed"],default=None)
    ap.add_argument("--force-edge",choices=["frontdoor_waf_cdn","cdn_then_waf","waf_then_cdn","waf_only","frontdoor_only","private_only"],default=None)
    ap.add_argument("--force-resilience",choices=["single_region_single_az","single_region_multi_az","multi_region_active_passive","multi_region_active_active","backup_only"],default=None)
    ap.add_argument("--force-data",choices=["transactional","enterprise_sql","postgres","event_driven","analytics","ai"],default=None)
    ap.add_argument("--force-network",choices=["hub_spoke","single_vnet","virtual_wan"],default=None)
    ap.add_argument("--output-dir",type=Path,default=Path("generated_dgf_docs"))
    args=ap.parse_args()
    args.output_dir.mkdir(parents=True,exist_ok=True)
    created=[]
    for i in range(args.count):
        overrides={"compute_profile":args.force_compute,"edge_pattern":args.force_edge,"resilience_profile":args.force_resilience,"data_profile":args.force_data,"network_profile":args.force_network}
        overrides={k:v for k,v in overrides.items() if v is not None}
        cdir, docs = build_case_pack(args.seed+i,args.difficulty,args.project_type,args.output_dir,args.conflict_rate,args.architecture_inconsistency_rate,overrides)
        created.append(str(cdir))
    print(json.dumps({"created_cases":created},ensure_ascii=False,indent=2))

if __name__=="__main__":
    main()
