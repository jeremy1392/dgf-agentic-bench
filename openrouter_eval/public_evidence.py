from __future__ import annotations
import csv, json, re
from pathlib import Path
from typing import Any
from docx import Document

DENY_NAMES = {"99_hidden_ground_truth.json", "tool_trace.jsonl"}


class PublicEvidenceReader:
    def __init__(self, case_dir: Path, gate: str, phase: str | None = None, max_chars: int = 40000, max_rows: int = 200):
        self.case_dir = Path(case_dir).resolve()
        self.gate = gate
        self.phase = phase
        self.max_chars = max_chars
        self.max_rows = max_rows
        self.graph = json.loads((self.case_dir / "02_evidence_graph.json").read_text(encoding="utf-8"))
        vis = self.case_dir / "05_phase_visibility.json"
        self.phase_visibility = json.loads(vis.read_text(encoding="utf-8")) if vis.exists() else {}
        self.nodes = {n["evidence_id"]: n for n in self.graph.get("nodes", [])}

    def _phase_ok(self, evidence_id: str) -> bool:
        if not self.phase:
            return True
        return evidence_id in set(self.phase_visibility.get(self.phase, []))

    def list_evidence(self, consumer_gate: str | None = None, include_unavailable: bool = True) -> dict[str, Any]:
        gate = consumer_gate or self.gate
        rows = []
        for n in self.nodes.values():
            if gate not in n.get("consumers", []):
                continue
            if not self._phase_ok(n["evidence_id"]):
                continue
            status = n.get("public_status", "AVAILABLE")
            if not include_unavailable and status != "AVAILABLE":
                continue
            rows.append({
                "evidence_id": n["evidence_id"],
                "path": n.get("path"),
                "public_status": status,
                "authoritative": bool(n.get("authoritative")),
                "version": n.get("version"),
                "age_days": n.get("age_days"),
                "consumers": n.get("consumers", []),
            })
        return {"gate": gate, "phase": self.phase, "evidence": rows}

    def _safe_path(self, node: dict[str, Any]) -> Path:
        rel = Path(str(node.get("path") or ""))
        if not rel or rel.name in DENY_NAMES or "hidden" in rel.name.lower():
            raise PermissionError("This file is not public evidence")
        p = (self.case_dir / rel).resolve()
        if self.case_dir not in p.parents and p != self.case_dir:
            raise PermissionError("Path escapes case directory")
        return p

    def read_evidence(self, evidence_id: str) -> dict[str, Any]:
        node = self.nodes.get(evidence_id)
        if not node:
            return {"status": "UNKNOWN_EVIDENCE", "evidence_id": evidence_id}
        if self.gate not in node.get("consumers", []):
            return {"status": "NOT_IN_GATE_SCOPE", "evidence_id": evidence_id, "gate": self.gate}
        if not self._phase_ok(evidence_id):
            return {"status": "NOT_AVAILABLE_IN_PHASE", "evidence_id": evidence_id, "phase": self.phase}
        if node.get("public_status") != "AVAILABLE":
            return {"status": node.get("public_status"), "evidence_id": evidence_id, "authoritative": bool(node.get("authoritative"))}
        p = self._safe_path(node)
        if not p.exists():
            return {"status": "MISSING_FILE", "evidence_id": evidence_id}
        suffix = p.suffix.lower()
        content: Any
        if suffix == ".json":
            content = json.loads(p.read_text(encoding="utf-8"))
        elif suffix == ".csv":
            with p.open(newline="", encoding="utf-8-sig") as f:
                content = list(csv.DictReader(f))[:self.max_rows]
        elif suffix == ".docx":
            doc = Document(p)
            blocks = []
            for par in doc.paragraphs:
                if par.text.strip():
                    blocks.append(par.text.strip())
            for t in doc.tables:
                for row in t.rows:
                    blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
            content = "\n".join(blocks)[:self.max_chars]
        elif suffix == ".svg":
            raw = p.read_text(encoding="utf-8", errors="replace")
            labels = re.findall(r"<text[^>]*>(.*?)</text>", raw, flags=re.I | re.S)
            labels = [re.sub(r"<[^>]+>", " ", x) for x in labels]
            labels = [re.sub(r"\s+", " ", x).strip() for x in labels if x.strip()]
            content = {"note": "Textual extraction from SVG; use authoritative resource tools for configuration facts.", "labels": labels[:300]}
        else:
            content = p.read_text(encoding="utf-8", errors="replace")[:self.max_chars]
        return {
            "status": "OK",
            "evidence_id": evidence_id,
            "authoritative": bool(node.get("authoritative")),
            "version": node.get("version"),
            "age_days": node.get("age_days"),
            "path": str(node.get("path")),
            "content": content,
        }
